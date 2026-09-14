from curl_cffi import requests as curl_req
import json
import re
import requests as req
from typing import List, Optional, Dict, Any, Callable
from bs4 import BeautifulSoup as bs 
import time, random

class PeterpanCrawler:
    
    def __init__(self):
        self.host = "https://www.peterpanz.com"
        self.search_url = "https://api.peterpanz.com/houses/area/pc"
        self.detail_url = lambda id: f"https://www.peterpanz.com/house/{id}"
        self.timeout_sec = 20
        
        self.HEADERS = {
            "accept": "application/json, text/plain, */*",
            "accept-encoding": "gzip, deflate, br, zstd",
            "accept-language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
            "content-type": "application/json",
            "origin": "https://www.peterpanz.com",
            "priority": "u=1, i",
            "referer": "https://www.peterpanz.com/",
            "sec-ch-ua": '"Not(A:Brand";v="8", "Chromium";v="144", "Google Chrome";v="144"',
            "sec-ch-ua-mobile": "?0",
            "sec-ch-ua-platform": '"Windows"',
            "sec-fetch-dest": "empty",
            "sec-fetch-mode": "cors",
            "sec-fetch-site": "same-site",
            "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/144.0.0.0 Safari/537.36",
            "x-identifier-id": "8348c629a0d191b6446256f97006e059",
            "x-peterpanz-os": "web",
            "x-peterpanz-page-id": "PAGE_UNKNOWN",
            "x-peterpanz-uidx": "undefined",
            "x-peterpanz-version": "3.60.0",
        }
        
        self.crawl_count = 0
        
    def _get_bbox(self, query: str) -> Dict[str, float]:
        url = "https://nominatim.openstreetmap.org/search"
        params = {
            "q": query,
            "format": "json",
            "limit": 1,
            "countrycodes": "kr",
        }
        headers = {"User-Agent": "myapp"}
        r = req.get(url, headers=headers, params=params, timeout=self.timeout_sec)
        data = r.json()[0]
        # boundingbox: [min_lat, max_lat, min_lon, max_lon]
        bb = data["boundingbox"]
        return {
            "lat_min": float(bb[0]),
            "lat_max": float(bb[1]),
            "lon_min": float(bb[2]),
            "lon_max": float(bb[3]),
        }
        
    def iter_search_page_ids(
        self,
        query: str = "서울특별시 중구 명동",
        building_type: str = "원/투룸",
        page_size: int = 50,
        latitude_min: Optional[float] = None,
        latitude_max: Optional[float] = None,
        longitude_min: Optional[float] = None,
        longitude_max: Optional[float] = None,
        check_deposit_min: Optional[int] = None,
        check_deposit_max: Optional[int] = None,
        check_month_min: Optional[int] = None,
        check_month_max: Optional[int] = None,
        check_real_size_min: Optional[float] = None,
        check_real_size_max: Optional[float] = None,
        check_price_min: Optional[int] = None,
        check_price_max: Optional[int] = None,
        contract_types: Optional[List[str]] = None,
        building_types: Optional[List[str]] = None,
        is_cancelled: Optional[Callable[[], bool]] = None,
    ):
        """지역 검색 결과를 페이지 단위 ID 리스트로 순회합니다."""
        has_bbox = all(
            x is not None for x in (latitude_min, latitude_max, longitude_min, longitude_max)
        )
        if has_bbox:
            bbox = {
                "lat_min": latitude_min,
                "lat_max": latitude_max,
                "lon_min": longitude_min,
                "lon_max": longitude_max,
            }
        else:
            bbox = self._get_bbox(query)
        parts: List[str] = []

        lat_min = latitude_min if latitude_min is not None else (bbox["lat_min"] - 0.02)
        lat_max = latitude_max if latitude_max is not None else (bbox["lat_max"] + 0.02)
        lon_min = longitude_min if longitude_min is not None else (bbox["lon_min"] - 0.025)
        lon_max = longitude_max if longitude_max is not None else (bbox["lon_max"] + 0.025)
        parts.append(f"latitude:{lat_min}~{lat_max}")
        parts.append(f"longitude:{lon_min}~{lon_max}")

        if check_deposit_min is not None or check_deposit_max is not None:
            mn = check_deposit_min if check_deposit_min is not None else 999
            mx = check_deposit_max if check_deposit_max is not None else 999
            parts.append(f"checkDeposit:{mn}~{mx}")
        if check_month_min is not None or check_month_max is not None:
            mn = check_month_min if check_month_min is not None else 999
            mx = check_month_max if check_month_max is not None else 999
            parts.append(f"checkMonth:{mn}~{mx}")
        if check_real_size_min is not None or check_real_size_max is not None:
            mn = check_real_size_min if check_real_size_min is not None else 999
            mx = check_real_size_max if check_real_size_max is not None else 999
            parts.append(f"checkRealSize:{mn}~{mx}")
        if check_price_min is not None or check_price_max is not None:
            mn = check_price_min if check_price_min is not None else 999
            mx = check_price_max if check_price_max is not None else 999
            parts.append(f"checkPrice:{mn}~{mx}")

        if contract_types:
            arr = json.dumps(contract_types, ensure_ascii=False)
            parts.append(f"contractType;{arr}")
        if building_types:
            arr = json.dumps(building_types, ensure_ascii=False)
            parts.append(f"buildingType;{arr}")
        elif building_type:
            parts.append(f'buildingType;["{building_type}"]')

        filter_str = "||".join(parts)
        parts = query.strip().rsplit(" ", 1)
        gungu = parts[0] if len(parts) > 1 else query
        dong = parts[1] if len(parts) > 1 else query

        seen: set = set()
        page_index = 1

        while True:
            if is_cancelled and is_cancelled():
                return
            
            params = {
                "zoomLevel": 15,
                "center": "",
                "adTarget": f'{{"type":"region","message":"{gungu}"}}',
                "dong": dong,
                "gungu": gungu,
                "filter": filter_str,
                "pageSize": page_size,
                "pageIndex": page_index,
                "order_id": 1772116818,
                "search": "",
                "filter_version": "5.1",
                "response_version": "5.3",
                "order_by": "random",
            }
            
            success = False
            
            for _ in range(3):
                if is_cancelled and is_cancelled():
                    return
                try:
                    r = curl_req.get(
                        self.search_url,
                        params=params,
                        headers=self.HEADERS,
                        impersonate="chrome",
                        timeout=self.timeout_sec,
                    )
                    r.raise_for_status()
                    success = True
                    break
                except Exception as e:
                    print(f"  실패: {e}")
                    if is_cancelled and is_cancelled():
                        return
                    time.sleep(random.uniform(0.2, 0.5))
                    continue
            if not success:
                raise Exception("검색 실패")
            
            houses = r.json().get("houses", {}).get("withoutFee", {}).get("image", [])
            if not houses:
                break
            page_ids = []
            for h in houses:
                hid = h.get("hidx")
                if not hid or hid in seen:
                    continue
                seen.add(hid)
                page_ids.append(hid)
            if not page_ids:
                break
            print(f"[피터팬] page {page_index}: {len(page_ids)}건")
            yield page_ids
            if len(page_ids) < page_size:
                break
            page_index += 1

    def search_get_ids(
        self,
        query: str = "서울특별시 중구 명동",
        building_type: str = "원/투룸",
        page_size: int = 50,
        latitude_min: Optional[float] = None,
        latitude_max: Optional[float] = None,
        longitude_min: Optional[float] = None,
        longitude_max: Optional[float] = None,
        check_deposit_min: Optional[int] = None,
        check_deposit_max: Optional[int] = None,
        check_month_min: Optional[int] = None,
        check_month_max: Optional[int] = None,
        check_real_size_min: Optional[float] = None,
        check_real_size_max: Optional[float] = None,
        check_price_min: Optional[int] = None,
        check_price_max: Optional[int] = None,
        contract_types: Optional[List[str]] = None,
        building_types: Optional[List[str]] = None,
    ) -> List[str]:
        """지역 검색어로 매물 ID 목록 조회. (기존 호환용)"""
        all_ids: List[str] = []
        for page_ids in self.iter_search_page_ids(
            query=query,
            building_type=building_type,
            page_size=page_size,
            latitude_min=latitude_min,
            latitude_max=latitude_max,
            longitude_min=longitude_min,
            longitude_max=longitude_max,
            check_deposit_min=check_deposit_min,
            check_deposit_max=check_deposit_max,
            check_month_min=check_month_min,
            check_month_max=check_month_max,
            check_real_size_min=check_real_size_min,
            check_real_size_max=check_real_size_max,
            check_price_min=check_price_min,
            check_price_max=check_price_max,
            contract_types=contract_types,
            building_types=building_types,
        ):
            all_ids.extend(page_ids)
        return all_ids
    
    def _parse_detail_table(self, soup: bs) -> Dict[str, Any]:
        """detail-table 영역에서 거래정보, 매물정보, 추가옵션, 시설정보, 상세설명 추출"""
        result = {"매물ID": None}
        detail_table = soup.find("div", class_="detail-table")
        if detail_table:
            sections = detail_table.find_all("h3", class_="detail-title")
            for h3 in sections:
                section_name = h3.get_text(strip=True)
                wrapper = h3.find_next_sibling("div", class_=re.compile("detail-table-wrapper"))
                if not wrapper:
                    continue
                rows = wrapper.find_all("div", class_="detail-table-row")
                for row in rows:
                    th = row.find("div", class_="detail-table-th")
                    td = row.find("div", class_="detail-table-td")
                    if th and td:
                        label = th.get_text(strip=True)
                        value = td.get_text(separator=" ", strip=True)
                        key = f"{section_name}_{label}" if section_name else label
                        result[key] = value

        # 추가옵션/시설정보: 태그/칩 형태일 수 있음
        for cls in ["additional-options", "facility-info", "option-tags"]:
            for el in soup.find_all(class_=re.compile(cls, re.I)):
                text = el.get_text(separator=", ", strip=True)
                if text:
                    result[cls] = result.get(cls, "") + (" " + text if result.get(cls) else text)

        # 상세설명: 별도 블록 (h3 다음 또는 특정 클래스)
        desc_title = soup.find("h3", string=re.compile("상세설명"))
        if desc_title:
            desc_block = desc_title.find_next_sibling()
            if desc_block:
                result["상세설명"] = desc_block.get_text(separator="\n", strip=True)
        else:
            desc_block = soup.find(class_=re.compile("detail-description|house-description|content"))
            if desc_block:
                result["상세설명"] = desc_block.get_text(separator="\n", strip=True)

        # 설명 박스(예: 최초 등록일)
        for li in soup.select("li.a-descriptionBox"):
            title_el = li.select_one(".a-descriptionBox__title")
            value_el = li.select_one(".a-descriptionBox__sub")
            if not title_el or not value_el:
                continue
            k = title_el.get_text(strip=True)
            v = value_el.get_text(strip=True)
            if k and v:
                result[k] = v

        # 위치정보: 주소 (house-address-type span.address)
        for addr_div in soup.find_all("div", class_="house-address-type"):
            addr_span = addr_div.find("span", class_="address")
            if addr_span:
                addr_type = addr_div.get("data-type", "")
                addr_val = addr_span.get_text(strip=True)
                if addr_type == "JIBUN":
                    result["주소"] = addr_val
                elif addr_type == "ROAD":
                    result["주소_도로명"] = addr_val
        if "주소" not in result and "주소_도로명" in result:
            result["주소"] = result["주소_도로명"]

        return result

    def _sanitize_excel(self, s: str) -> str:
        """엑셀 제어문자 제거 (IllegalCharacterError 방지)"""
        if not s:
            return ""
        return "".join(
            c if (ord(c) in (0x09, 0x0A, 0x0D) or 0x20 <= ord(c) <= 0x7E or ord(c) >= 0xA0) else " "
            for c in str(s)
        )

    def crawl_detail(self, id: str, is_cancelled: Optional[Callable[[], bool]] = None) -> Dict[str, Any]:
        """상세 페이지 크롤링 후 추출 데이터 반환"""
        for _ in range(3):
            if is_cancelled and is_cancelled():
                raise InterruptedError("중단 요청됨")
            try:
                r = curl_req.get(
                    self.detail_url(id),
                    headers=self.HEADERS,
                    impersonate="chrome",
                    timeout=self.timeout_sec,
                )
                r.raise_for_status()
                break
            except Exception as e:
                print(f"  실패: {e}")
                if is_cancelled and is_cancelled():
                    raise InterruptedError("중단 요청됨")
                time.sleep(random.uniform(0.2, 0.5))
                continue
                
        soup = bs(r.text, "html.parser")
        data = self._parse_detail_table(soup)
        data["매물ID"] = id
        data["URL"] = self.detail_url(id)
        
        self.crawl_count += 1
        print(f"[피터팬] {self.crawl_count}건: {data['매물정보_건축물용도']}, {data['주소']}")
        
        return data

    def crawl_details_to_excel(
        self,
        ids: List[str],
        output_path: str = "peterpan_detail.xlsx",
        on_item: Optional[Callable[[Dict[str, Any]], None]] = None,
    ) -> str:
        """여러 매물 상세를 크롤링해 엑셀에 저장. on_item(row) 콜백으로 각 건마다 호출."""
        try:
            import openpyxl
        except ImportError:
            raise ImportError("pip install openpyxl")

        rows = []
        for i, hid in enumerate(ids, 1):
            print(f"[{i}/{len(ids)}] {hid}")
            try:
                data = self.crawl_detail(str(hid))
                rows.append(data)
                if on_item:
                    on_item(data)
            except Exception as e:
                print(f"  실패: {e}")
                rows.append({"매물ID": hid, "오류": str(e)})
            finally:
                time.sleep(random.uniform(0.1, 0.3))

        return self.save_rows_to_excel(rows, output_path)

    def save_rows_to_excel(self, rows: List[Dict[str, Any]], output_path: str) -> str:
        try:
            import openpyxl
        except ImportError:
            raise ImportError("pip install openpyxl")
        if not rows:
            raise ValueError("저장할 데이터 없음")
        all_keys = []
        seen = set()
        for r in rows:
            for k in r.keys():
                if k not in seen:
                    seen.add(k)
                    all_keys.append(k)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "매물상세"
        ws.append([self._sanitize_excel(str(k)) for k in all_keys])
        for r in rows:
            ws.append([self._sanitize_excel(str(r.get(k, ""))) for k in all_keys])
        self._apply_hyperlinks(ws, all_keys)
        wb.save(output_path)
        return output_path

    def _apply_hyperlinks(self, ws, all_keys: List[str]) -> None:
        """URL/링크 컬럼을 엑셀 하이퍼링크 셀로 저장"""
        try:
            from openpyxl.styles import Font
        except Exception:
            return
        url_col_indices = []
        for i, key in enumerate(all_keys, 1):
            k = str(key).strip().lower()
            if ("url" in k) or ("link" in k) or ("링크" in str(key)):
                url_col_indices.append(i)
        if not url_col_indices:
            return
        link_font = Font(color="0563C1", underline="single")
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in url_col_indices:
                cell = ws.cell(row=row_idx, column=col_idx)
                value = str(cell.value or "").strip()
                if value.lower().startswith(("http://", "https://")):
                    cell.hyperlink = value
                    cell.font = link_font


if __name__ == "__main__":
    crawler = PeterpanCrawler()
    # 서울특별시 + 전 매물유형 + 전 거래유형 전체 수집
    ids = crawler.search_get_ids(
        query="서울특별시",
        building_type="",
        contract_types=["월세", "전세", "단기임대", "매매"],
        building_types=["원/투룸", "빌라/주택", "오피스텔", "아파트"],
    )
    print(f"[서울특별시] 매물 {len(ids)}개")
    if ids:
        path = crawler.crawl_details_to_excel(ids, "peterpan_서울특별시_전체유형_전체거래.xlsx")
        print(f"저장: {path}")