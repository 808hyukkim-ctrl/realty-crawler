# -*- coding: utf-8 -*-
"""
당근알바(jobs.daangn.com) 수집 - 데이터 모듈.

사실(2026-10-12 확인)
- GraphQL: POST https://jobs.kr.karrotmarket.com/graphql  {id:<persisted id>, query:"", variables}  (로그인 불필요)
    routeSearchResultsQuery      b479f2c2… : {where:{searchKeyword, region:<동 _id>, regionRange:"RANGE_4"}, first, filter, isStaffCuration:false, after, othersFilter, priorityFilter}
    SearchResultsPaginationQuery daf6296e… : {where, first, filter, after}
    routeJobPostByPublicIdQuery  53f9195a… : {publicId}  → permalinkPairByPublicId.daangnPermalink.model (content, 주소, 조회수, 지원자수…)
    regionCascadeSelectColumnQuery 3b922919… : {parentRegionId:<int|null>} → regionsByParentRegionId (시도→시군구→동)
    keywordResultsQuery          e5eb7d7e… : {query, first:20} → searchRegions (동 이름 검색)
- where.region 은 동(depth 3) _id 만 됨(구·시도 id 는 0건). RANGE_4 가 가장 넓음(이웃 동 포함). 한 쪽 최대 50건쯤(first=50 → 39).
- filter.workPeriod LESS_THAN_A_MONTH = 단기, MORE_THAN_A_MONTH = 1개월 이상. employmentType STAFF 정직원 / PART_TIME_JOB 알바.
- 웹 주소: https://jobs.daangn.com/job-posts/<slug>-<publicId>  (slug 없이 publicId 만으로도 열림)
"""
from __future__ import annotations

import json
import os
import random
import time
from datetime import date, datetime, timedelta
from typing import Any, Callable, Dict, List, Optional, Set, Tuple

import requests

GQL = "https://jobs.kr.karrotmarket.com/graphql"
Q_SEARCH = "b479f2c233edf0d921b4b954b8e1a0ff"
Q_PAGE = "daf6296ed9995299770c6ade90b9d80c"
Q_DETAIL = "53f9195a00c59634ffdabca4178f7581"
Q_REGION_CHILDREN = "3b92291962883ef85597e461e8908e89"
Q_REGION_SEARCH = "e5eb7d7e6ca3a71821db69609e66d47d"
POST_URL = "https://jobs.daangn.com/job-posts/{pid}"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
    "Accept": "application/json",
    "content-type": "application/json",
    "Origin": "https://jobs.daangn.com",
    "Referer": "https://jobs.daangn.com/",
    "x-graphql-client-name": "jobs-website",
}

SALARY_KO = {"HOURLY": "시급", "DAILY": "일급", "MONTHLY": "월급", "YEARLY": "연봉", "PER_CASE": "건당"}
EMPLOY_KO = {"PART_TIME_JOB": "알바", "STAFF": "정직원"}
PERIOD_KO = {"LESS_THAN_A_MONTH": "단기(1개월 미만)", "MORE_THAN_A_MONTH": "1개월 이상"}
JOB_TASK_KO = {
    "BAKING": "베이킹", "BEAUTY_CARE": "미용/뷰티", "BEVERAGE_MAKING": "음료 제조", "CAFE_BAKERY": "카페/베이커리",
    "CHILD_CARE": "아이돌봄", "CLEANING": "업체청소", "CONSTRUCTION_FIELD": "건설/현장", "CONVENIENCE_STORE": "편의점",
    "ETC": "기타", "HOUSEWORK": "가사/청소/정리", "INTERIOR": "인테리어", "KIDS_PICK_UP": "등하원도우미",
    "KITCHEN_ASSISTANCE": "주방보조/설거지", "KITCHEN_COOK": "주방장/조리사", "LEAFLET_DISTRIBUTION": "전단지",
    "LIGHT_WORK": "심부름/소일거리", "LOGISTICS_PACKAGING": "물류/포장", "MANUFACTURING": "생산/제조",
    "MOVING_ASSISTANCE": "짐 옮기기", "NURSING": "돌봄/시팅", "OFFICE_ASSISTANCE": "사무보조", "PET_CARE": "반려동물 돌봄",
    "REPAIR_INSTALLATION": "수리/설치", "SENIOR_CARE": "어르신 돌봄", "SERVING": "서빙", "SHOP_MANAGEMENT": "매장관리/판매",
    "TUTORING": "학원/과외/레슨",
}
DAY_KO = {"MON": "월", "TUE": "화", "WED": "수", "THU": "목", "FRI": "금", "SAT": "토", "SUN": "일"}

# 사이트가 안 될 때 쓰는 시도 목록 (2026-10-12)
BUILTIN_SIDOS: List[Dict[str, Any]] = [
    {"_id": 1873, "name": "강원특별자치도"},
    {"_id": 1256, "name": "경기도"},
    {"_id": 3464, "name": "경상남도"},
    {"_id": 3093, "name": "경상북도"},
    {"_id": 675, "name": "대구광역시"},
    {"_id": 1094, "name": "대전광역시"},
    {"_id": 451, "name": "부산광역시"},
    {"_id": 1, "name": "서울특별시"},
    {"_id": 1241, "name": "세종특별자치시"},
    {"_id": 1179, "name": "울산광역시"},
    {"_id": 825, "name": "인천광역시"},
    {"_id": 13371, "name": "전남광주통합특별시"},
    {"_id": 2488, "name": "전북특별자치도"},
    {"_id": 3811, "name": "제주특별자치도"},
    {"_id": 2261, "name": "충청남도"},
    {"_id": 2092, "name": "충청북도"},
]

# 시도별 시군구 목록 (2026-10-13 기준, 사이트가 안 될 때도 바로 고를 수 있게)
BUILTIN_GUS: Dict[str, List[Dict[str, Any]]] = {
    "1873": [{"_id": 1927, "name": "강릉시"}, {"_id": 2078, "name": "고성군"}, {"_id": 1949, "name": "동해시"}, {"_id": 1978, "name": "삼척시"}, {"_id": 1969, "name": "속초시"}, {"_id": 2064, "name": "양구군"}, {"_id": 2085, "name": "양양군"}, {"_id": 2013, "name": "영월군"}, {"_id": 1901, "name": "원주시"}, {"_id": 2070, "name": "인제군"}, {"_id": 2034, "name": "정선군"}, {"_id": 2045, "name": "철원군"}, {"_id": 1875, "name": "춘천시"}, {"_id": 1960, "name": "태백시"}, {"_id": 2024, "name": "평창군"}, {"_id": 1992, "name": "홍천군"}, {"_id": 2058, "name": "화천군"}, {"_id": 2003, "name": "횡성군"}],
    "1256": [{"_id": 1853, "name": "가평군"}, {"_id": 1529, "name": "고양시 덕양구"}, {"_id": 1549, "name": "고양시 일산동구"}, {"_id": 1561, "name": "고양시 일산서구"}, {"_id": 1571, "name": "과천시"}, {"_id": 1447, "name": "광명시"}, {"_id": 1791, "name": "광주시"}, {"_id": 1578, "name": "구리시"}, {"_id": 1631, "name": "군포시"}, {"_id": 1751, "name": "김포시"}, {"_id": 1587, "name": "남양주시"}, {"_id": 1491, "name": "동두천시"}, {"_id": 1429, "name": "부천시 소사구"}, {"_id": 1439, "name": "부천시 오정구"}, {"_id": 1408, "name": "부천시 원미구"}, {"_id": 1334, "name": "성남시 분당구"}, {"_id": 1305, "name": "성남시 수정구"}, {"_id": 1322, "name": "성남시 중원구"}, {"_id": 1270, "name": "수원시 권선구"}, {"_id": 1293, "name": "수원시 영통구"}, {"_id": 1259, "name": "수원시 장안구"}, {"_id": 1282, "name": "수원시 팔달구"}, {"_id": 1613, "name": "시흥시"}, {"_id": 1515, "name": "안산시 단원구"}, {"_id": 1501, "name": "안산시 상록구"}, {"_id": 1735, "name": "안성시"}, {"_id": 1389, "name": "안양시 동안구"}, {"_id": 1374, "name": "안양시 만안구"}, {"_id": 1802, "name": "양주시"}, {"_id": 1860, "name": "양평군"}, {"_id": 1829, "name": "여주시"}, {"_id": 1842, "name": "연천군"}, {"_id": 1606, "name": "오산시"}, {"_id": 1676, "name": "용인시 기흥구"}, {"_id": 1688, "name": "용인시 수지구"}, {"_id": 1664, "name": "용인시 처인구"}, {"_id": 1643, "name": "의왕시"}, {"_id": 1357, "name": "의정부시"}, {"_id": 1720, "name": "이천시"}, {"_id": 1698, "name": "파주시"}, {"_id": 1466, "name": "평택시"}, {"_id": 1814, "name": "포천시"}, {"_id": 1650, "name": "하남시"}, {"_id": 13301, "name": "화성시 동탄구"}, {"_id": 13298, "name": "화성시 만세구"}, {"_id": 13300, "name": "화성시 병점구"}, {"_id": 13299, "name": "화성시 효행구"}],
    "3464": [{"_id": 3638, "name": "거제시"}, {"_id": 3780, "name": "거창군"}, {"_id": 3716, "name": "고성군"}, {"_id": 3599, "name": "김해시"}, {"_id": 3731, "name": "남해군"}, {"_id": 3620, "name": "밀양시"}, {"_id": 3582, "name": "사천시"}, {"_id": 3756, "name": "산청군"}, {"_id": 3661, "name": "양산시"}, {"_id": 3676, "name": "의령군"}, {"_id": 3533, "name": "진주시"}, {"_id": 3701, "name": "창녕군"}, {"_id": 3483, "name": "창원시 마산합포구"}, {"_id": 3503, "name": "창원시 마산회원구"}, {"_id": 3475, "name": "창원시 성산구"}, {"_id": 3466, "name": "창원시 의창구"}, {"_id": 3517, "name": "창원시 진해구"}, {"_id": 3566, "name": "통영시"}, {"_id": 3742, "name": "하동군"}, {"_id": 3690, "name": "함안군"}, {"_id": 3768, "name": "함양군"}, {"_id": 3793, "name": "합천군"}],
    "3093": [{"_id": 3310, "name": "경산시"}, {"_id": 3127, "name": "경주시"}, {"_id": 3394, "name": "고령군"}, {"_id": 3199, "name": "구미시"}, {"_id": 3151, "name": "김천시"}, {"_id": 3293, "name": "문경시"}, {"_id": 3436, "name": "봉화군"}, {"_id": 3264, "name": "상주시"}, {"_id": 3403, "name": "성주군"}, {"_id": 3174, "name": "안동시"}, {"_id": 3370, "name": "영덕군"}, {"_id": 3363, "name": "영양군"}, {"_id": 3227, "name": "영주시"}, {"_id": 3247, "name": "영천시"}, {"_id": 3423, "name": "예천군"}, {"_id": 3459, "name": "울릉군"}, {"_id": 3447, "name": "울진군"}, {"_id": 3335, "name": "의성군"}, {"_id": 3382, "name": "청도군"}, {"_id": 3354, "name": "청송군"}, {"_id": 3414, "name": "칠곡군"}, {"_id": 3095, "name": "포항시 남구"}, {"_id": 3110, "name": "포항시 북구"}],
    "675": [{"_id": 3326, "name": "군위군"}, {"_id": 728, "name": "남구"}, {"_id": 790, "name": "달서구"}, {"_id": 813, "name": "달성군"}, {"_id": 689, "name": "동구"}, {"_id": 742, "name": "북구"}, {"_id": 710, "name": "서구"}, {"_id": 766, "name": "수성구"}, {"_id": 676, "name": "중구"}],
    "1094": [{"_id": 1166, "name": "대덕구"}, {"_id": 1095, "name": "동구"}, {"_id": 1130, "name": "서구"}, {"_id": 1154, "name": "유성구"}, {"_id": 1112, "name": "중구"}],
    "451": [{"_id": 624, "name": "강서구"}, {"_id": 606, "name": "금정구"}, {"_id": 669, "name": "기장군"}, {"_id": 538, "name": "남구"}, {"_id": 476, "name": "동구"}, {"_id": 524, "name": "동래구"}, {"_id": 502, "name": "부산진구"}, {"_id": 556, "name": "북구"}, {"_id": 656, "name": "사상구"}, {"_id": 589, "name": "사하구"}, {"_id": 462, "name": "서구"}, {"_id": 645, "name": "수영구"}, {"_id": 632, "name": "연제구"}, {"_id": 490, "name": "영도구"}, {"_id": 452, "name": "중구"}, {"_id": 570, "name": "해운대구"}],
    "1": [{"_id": 381, "name": "강남구"}, {"_id": 432, "name": "강동구"}, {"_id": 140, "name": "강북구"}, {"_id": 257, "name": "강서구"}, {"_id": 340, "name": "관악구"}, {"_id": 71, "name": "광진구"}, {"_id": 278, "name": "구로구"}, {"_id": 294, "name": "금천구"}, {"_id": 169, "name": "노원구"}, {"_id": 154, "name": "도봉구"}, {"_id": 87, "name": "동대문구"}, {"_id": 324, "name": "동작구"}, {"_id": 221, "name": "마포구"}, {"_id": 206, "name": "서대문구"}, {"_id": 362, "name": "서초구"}, {"_id": 53, "name": "성동구"}, {"_id": 119, "name": "성북구"}, {"_id": 404, "name": "송파구"}, {"_id": 238, "name": "양천구"}, {"_id": 305, "name": "영등포구"}, {"_id": 36, "name": "용산구"}, {"_id": 189, "name": "은평구"}, {"_id": 2, "name": "종로구"}, {"_id": 20, "name": "중구"}, {"_id": 102, "name": "중랑구"}],
    "1241": [{"_id": 6412, "name": "가람동"}, {"_id": 4186, "name": "고운동"}, {"_id": 1247, "name": "금남면"}, {"_id": 6413, "name": "나성동"}, {"_id": 6414, "name": "다정동"}, {"_id": 4247, "name": "대평동"}, {"_id": 1254, "name": "도담동"}, {"_id": 6415, "name": "반곡동"}, {"_id": 4184, "name": "보람동"}, {"_id": 1246, "name": "부강면"}, {"_id": 12547, "name": "산울동"}, {"_id": 4204, "name": "새롬동"}, {"_id": 4248, "name": "소담동"}, {"_id": 1252, "name": "소정면"}, {"_id": 1255, "name": "아름동"}, {"_id": 6416, "name": "어진동"}, {"_id": 1244, "name": "연기면"}, {"_id": 1245, "name": "연동면"}, {"_id": 1249, "name": "연서면"}, {"_id": 1248, "name": "장군면"}, {"_id": 1251, "name": "전동면"}, {"_id": 1250, "name": "전의면"}, {"_id": 1243, "name": "조치원읍"}, {"_id": 4153, "name": "종촌동"}, {"_id": 10941, "name": "집현동"}, {"_id": 1253, "name": "한솔동"}, {"_id": 10677, "name": "해밀동"}],
    "1179": [{"_id": 1194, "name": "남구"}, {"_id": 1209, "name": "동구"}, {"_id": 1219, "name": "북구"}, {"_id": 1228, "name": "울주군"}, {"_id": 1180, "name": "중구"}],
    "825": [{"_id": 968, "name": "강화군"}, {"_id": 13370, "name": "검단구"}, {"_id": 933, "name": "계양구"}, {"_id": 890, "name": "남동구"}, {"_id": 854, "name": "미추홀구"}, {"_id": 910, "name": "부평구"}, {"_id": 946, "name": "서해구"}, {"_id": 876, "name": "연수구"}, {"_id": 13368, "name": "영종구"}, {"_id": 983, "name": "옹진군"}, {"_id": 13369, "name": "제물포구"}],
    "13371": [{"_id": 2957, "name": "강진군"}, {"_id": 2898, "name": "고흥군"}, {"_id": 2877, "name": "곡성군"}, {"_id": 1072, "name": "광산구"}, {"_id": 2851, "name": "광양시"}, {"_id": 2889, "name": "구례군"}, {"_id": 2830, "name": "나주시"}, {"_id": 1027, "name": "남구"}, {"_id": 2864, "name": "담양군"}, {"_id": 994, "name": "동구"}, {"_id": 2749, "name": "목포시"}, {"_id": 2997, "name": "무안군"}, {"_id": 2917, "name": "보성군"}, {"_id": 1044, "name": "북구"}, {"_id": 1008, "name": "서구"}, {"_id": 2804, "name": "순천시"}, {"_id": 3068, "name": "신안군"}, {"_id": 2773, "name": "여수시"}, {"_id": 3018, "name": "영광군"}, {"_id": 2984, "name": "영암군"}, {"_id": 3043, "name": "완도군"}, {"_id": 3031, "name": "장성군"}, {"_id": 2946, "name": "장흥군"}, {"_id": 3058, "name": "진도군"}, {"_id": 3008, "name": "함평군"}, {"_id": 2969, "name": "해남군"}, {"_id": 2931, "name": "화순군"}],
    "2488": [{"_id": 2719, "name": "고창군"}, {"_id": 2526, "name": "군산시"}, {"_id": 2633, "name": "김제시"}, {"_id": 2609, "name": "남원시"}, {"_id": 2679, "name": "무주군"}, {"_id": 2734, "name": "부안군"}, {"_id": 2707, "name": "순창군"}, {"_id": 2653, "name": "완주군"}, {"_id": 2554, "name": "익산시"}, {"_id": 2694, "name": "임실군"}, {"_id": 2686, "name": "장수군"}, {"_id": 2509, "name": "전주시 덕진구"}, {"_id": 2490, "name": "전주시 완산구"}, {"_id": 2585, "name": "정읍시"}, {"_id": 2667, "name": "진안군"}],
    "3811": [{"_id": 3839, "name": "서귀포시"}, {"_id": 3812, "name": "제주시"}],
    "2261": [{"_id": 2381, "name": "계룡시"}, {"_id": 2295, "name": "공주시"}, {"_id": 2401, "name": "금산군"}, {"_id": 2365, "name": "논산시"}, {"_id": 2386, "name": "당진시"}, {"_id": 2312, "name": "보령시"}, {"_id": 2412, "name": "부여군"}, {"_id": 2349, "name": "서산시"}, {"_id": 2429, "name": "서천군"}, {"_id": 2331, "name": "아산시"}, {"_id": 2466, "name": "예산군"}, {"_id": 2263, "name": "천안시 동남구"}, {"_id": 2281, "name": "천안시 서북구"}, {"_id": 2443, "name": "청양군"}, {"_id": 2479, "name": "태안군"}, {"_id": 2454, "name": "홍성군"}],
    "2092": [{"_id": 2230, "name": "괴산군"}, {"_id": 2252, "name": "단양군"}, {"_id": 2185, "name": "보은군"}, {"_id": 2207, "name": "영동군"}, {"_id": 2197, "name": "옥천군"}, {"_id": 2242, "name": "음성군"}, {"_id": 2167, "name": "제천시"}, {"_id": 2219, "name": "증평군"}, {"_id": 2222, "name": "진천군"}, {"_id": 2094, "name": "청주시 상당구"}, {"_id": 2108, "name": "청주시 서원구"}, {"_id": 2132, "name": "청주시 청원구"}, {"_id": 2120, "name": "청주시 흥덕구"}, {"_id": 2141, "name": "충주시"}],
}

EXCEL_COLUMNS: List[str] = [
    "번호", "제목", "회사명", "지역", "주소", "급여", "고용형태", "근무기간", "업무", "근무요일/날짜", "근무시간",
    "내용", "전화연락", "조회수", "관심", "지원자", "등록일시", "마감", "검색어", "당근링크",
]


def to_kst(iso: Optional[str]) -> Optional[datetime]:
    if not iso:
        return None
    try:
        return datetime.strptime(str(iso)[:19], "%Y-%m-%dT%H:%M:%S") + timedelta(hours=9)
    except Exception:
        return None


def salary_text(jp: Dict[str, Any]) -> str:
    st = jp.get("salaryType") or ""
    s = jp.get("salary")
    lab = SALARY_KO.get(st, st)
    if s is None:
        return lab
    try:
        n = int(s)
    except Exception:
        return f"{lab} {s}"
    if st in ("MONTHLY", "YEARLY"):
        return f"{lab} {n // 10000:,}만원"
    txt = f"{lab} {n:,}원"
    if jp.get("salaryPayMethod") == "WORK_DAY_PAYMENT":
        txt += " (당일지급)"
    return txt


def workdays_text(jp: Dict[str, Any]) -> str:
    days = [DAY_KO.get(d, d) for d in (jp.get("workDays") or [])]
    dates = jp.get("workDates") or []
    parts = []
    if days:
        parts.append(",".join(days))
    if dates:
        ds = sorted(dates)
        parts.append(ds[0] if len(ds) == 1 else f"{ds[0]}~{ds[-1]} ({len(ds)}일)")
    if jp.get("isWorkDaysNegotiable"):
        parts.append("요일협의")
    return " ".join(parts)


def worktime_text(jp: Dict[str, Any]) -> str:
    a, b = jp.get("workTimeStart"), jp.get("workTimeEnd")
    t = f"{a}~{b}" if a and b else (a or b or "")
    if jp.get("isWorkTimeNegotiable"):
        t = (t + " 시간협의").strip()
    return t


def normalize(jp: Dict[str, Any], keyword: str = "", detail: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    d = detail or {}
    pid = ((jp.get("daangnPermalink") or {}).get("publicId")) or ((d.get("daangnPermalink") or {}).get("publicId")) or ""
    region = ((jp.get("workplaceRegion") or {}).get("name3Region") or {}).get("name") or (jp.get("workplaceRegion") or {}).get("name") or ""
    pub = to_kst(jp.get("publishedAt") or d.get("publishedAt"))
    closed = bool(jp.get("closed") or d.get("closed"))
    tasks = ", ".join(JOB_TASK_KO.get(t, t) for t in (d.get("jobTasks") or jp.get("jobTasks") or []) if t != "ETC")
    addr = d.get("workplaceRoadAddress") or d.get("workplaceAddress") or ""
    if d.get("workplaceAddress") and d.get("workplaceRoadAddress") and d["workplaceAddress"] != d["workplaceRoadAddress"]:
        addr = f"{d['workplaceRoadAddress']} ({d['workplaceAddress']})"
    app = d.get("applicationCountInfo") or {}
    return {
        "번호": pid,
        "제목": (jp.get("title") or d.get("title") or "").strip(),
        "회사명": jp.get("companyName") or d.get("companyName") or "",
        "지역": region,
        "주소": addr,
        "급여": salary_text(jp if jp.get("salaryType") else d),
        "고용형태": EMPLOY_KO.get(jp.get("employmentType") or d.get("employmentType") or "", jp.get("employmentType") or ""),
        "근무기간": PERIOD_KO.get(d.get("workPeriod") or jp.get("workPeriod") or "", ""),
        "업무": tasks,
        "근무요일/날짜": workdays_text(jp),
        "근무시간": worktime_text(jp),
        "내용": (d.get("content") or "").strip(),
        "전화연락": "가능" if d.get("isPhoneContactable") else ("" if not d else "채팅만"),
        "조회수": d.get("viewCount") if d else "",
        "관심": d.get("watchCount") if d else "",
        "지원자": app.get("totalCount") if d else "",
        "등록일시": pub.strftime("%Y-%m-%d %H:%M") if pub else "",
        "마감": "마감" if closed else "",
        "검색어": keyword,
        "당근링크": POST_URL.format(pid=pid) if pid else "",
        "_created": pub,
        "_id": jp.get("_id") or d.get("_id"),
    }


class DaangnJobsClient:
    def __init__(self, log: Optional[Callable[[str], None]] = None, delay: Tuple[float, float] = (0.3, 0.8)):
        self.log = log or (lambda s: None)
        self.delay = delay
        self.s = requests.Session()
        self.s.headers.update(HEADERS)
        self._n = 0

    def _sleep(self) -> None:
        if self._n and self.delay[1] > 0:
            time.sleep(random.uniform(*self.delay))
        self._n += 1

    def gql(self, qid: str, variables: Dict[str, Any], tries: int = 3) -> Dict[str, Any]:
        last: Any = None
        for i in range(tries):
            try:
                self._sleep()
                r = self.s.post(GQL, data=json.dumps({"id": qid, "query": "", "variables": variables}), timeout=30)
                if r.status_code == 429 or r.status_code >= 500:
                    raise RuntimeError(f"HTTP {r.status_code}")
                j = r.json()
                if j.get("errors") and not j.get("data"):
                    raise RuntimeError("당근알바 응답 오류: " + str(j["errors"][0].get("message", ""))[:200])
                return j.get("data") or {}
            except Exception as e:  # noqa: BLE001
                last = e
                if i < tries - 1:
                    wait = 3 * (i + 1)
                    self.log(f"  재시도 {i + 1}/{tries - 1} ({e}) {wait}초 대기")
                    time.sleep(wait)
        raise RuntimeError(f"당근알바 접속 실패: {last}")

    # ---- 지역 ----
    def region_children(self, parent_id: Optional[int]) -> List[Dict[str, Any]]:
        d = self.gql(Q_REGION_CHILDREN, {"parentRegionId": parent_id})
        return [{"_id": r["_id"], "name": r.get("name") or ""} for r in (d.get("regionsByParentRegionId") or [])]

    def sidos(self) -> List[Dict[str, Any]]:
        try:
            out = self.region_children(None)
            return out or list(BUILTIN_SIDOS)
        except Exception:
            return list(BUILTIN_SIDOS)

    def search_regions(self, query: str) -> List[Dict[str, Any]]:
        d = self.gql(Q_REGION_SEARCH, {"query": query, "first": 20})
        out = []
        for e in ((d.get("searchRegions") or {}).get("edges") or []):
            n = e.get("node") or {}
            out.append({"_id": n.get("_id"), "name": n.get("name"), "name1": n.get("name1"), "name2": n.get("name2"), "name3": n.get("name3")})
        return out

    # ---- 검색 ----
    def search(self, region_id: int, keyword: str = "", short_term: bool = False, after: Optional[str] = None,
               first: int = 50) -> Tuple[List[Dict[str, Any]], Optional[str], bool]:
        where = {"searchKeyword": keyword or "", "region": int(region_id), "regionRange": "RANGE_4"}
        flt: Dict[str, Any] = {}
        if short_term:
            flt["workPeriod"] = "LESS_THAN_A_MONTH"
        if after:
            d = self.gql(Q_PAGE, {"where": where, "first": first, "filter": flt, "after": after})
        else:
            d = self.gql(Q_SEARCH, {"where": where, "first": first, "filter": flt, "isStaffCuration": False,
                                     "after": None, "othersFilter": flt, "priorityFilter": flt})
        feed = d.get("searchedFeed") or {}
        posts = []
        for e in feed.get("edges") or []:
            jp = (e.get("node") or {}).get("jobPost")
            if jp:
                posts.append(jp)
        pi = feed.get("pageInfo") or {}
        return posts, pi.get("endCursor"), bool(pi.get("hasNextPage"))

    def detail(self, public_id: str) -> Dict[str, Any]:
        d = self.gql(Q_DETAIL, {"publicId": public_id})
        pair = d.get("permalinkPairByPublicId") or {}
        model = (pair.get("daangnPermalink") or {}).get("model") or (pair.get("permalink") or {}).get("model") or {}
        if model.get("__typename") not in (None, "JobPost"):
            return {}
        return model


def _contains_any(text: str, words: List[str]) -> bool:
    t = (text or "").lower()
    return any(w.lower() in t for w in words if w)


def collect(
    client: DaangnJobsClient,
    region_ids: List[int],
    keywords: List[str],
    start: date,
    end: date,
    max_count: int = 500,
    short_term: bool = False,
    with_detail: bool = True,
    include_words: Optional[List[str]] = None,
    exclude_words: Optional[List[str]] = None,
    skip_closed: bool = True,
    seen_ids: Optional[Set[str]] = None,
    stop: Optional[Callable[[], bool]] = None,
    log: Optional[Callable[[str], None]] = None,
    progress: Optional[Callable[[int, int], None]] = None,
    region_names: Optional[Dict[int, str]] = None,
    max_pages: int = 10,
) -> List[Dict[str, Any]]:
    """동 id 목록 × 검색어 별로 당근알바를 검색해 중복 없이 모은다(등록일 기간, 최대 건수)."""
    log = log or (lambda s: None)
    stop = stop or (lambda: False)
    include_words = [w for w in (include_words or []) if w]
    exclude_words = [w for w in (exclude_words or []) if w]
    seen_ids = seen_ids or set()
    kws: List[str] = [k for k in (keywords or []) if k] or [""]
    names = region_names or {}

    got: Dict[str, Dict[str, Any]] = {}
    order: List[str] = []
    kw_of: Dict[str, List[str]] = {}
    for rid in region_ids:
        if stop():
            break
        for kw in kws:
            if stop():
                break
            after = None
            pages = 0
            older_streak = 0
            label = f"{names.get(rid, rid)}" + (f" '{kw}'" if kw else "")
            n_before = len(order)
            while not stop():
                try:
                    posts, after, more = client.search(rid, kw, short_term, after)
                except Exception as e:  # noqa: BLE001
                    log(f"  [{label}] 검색 실패: {e}")
                    break
                pages += 1
                for jp in posts:
                    pub = to_kst(jp.get("publishedAt"))
                    d = pub.date() if pub else None
                    if d and d < start:
                        older_streak += 1
                        continue
                    if d and d > end:
                        continue
                    older_streak = 0
                    pid = jp.get("_id") or ""
                    if not pid:
                        continue
                    if pid in got:
                        if kw and kw not in kw_of[pid]:
                            kw_of[pid].append(kw)
                        continue
                    if pid in seen_ids:
                        continue
                    if skip_closed and jp.get("closed"):
                        continue
                    n = normalize(jp, kw)
                    hay = n["제목"] + " " + n["회사명"]
                    if exclude_words and _contains_any(hay, exclude_words):
                        continue
                    got[pid] = n
                    kw_of[pid] = [kw] if kw else []
                    order.append(pid)
                    if max_count and len(order) >= max_count:
                        break
                if progress:
                    progress(len(order), max_count or 0)
                if (max_count and len(order) >= max_count) or not more or not after or pages >= max_pages or older_streak >= 20:
                    break
            log(f"  [{label}] +{len(order) - n_before}건 (누적 {len(order)})")
            if max_count and len(order) >= max_count:
                break
        if max_count and len(order) >= max_count:
            log(f"최대 건수 {max_count}건에 도달해 멈춥니다.")
            break

    rows = [got[p] for p in order]
    if with_detail and rows:
        log(f"상세(내용·주소·조회수) 가져오는 중... {len(rows)}건")
        for i, r in enumerate(rows, 1):
            if stop():
                break
            try:
                dt = client.detail(r["번호"]) if r.get("번호") else {}
            except Exception as e:  # noqa: BLE001
                log(f"  상세 실패 {r.get('번호')}: {e}")
                dt = {}
            if dt:
                jp_stub = {k: r.get(k) for k in ()}  # noqa: F841
                merged = normalize({**dt, "daangnPermalink": {"publicId": r["번호"]}, "workplaceRegion": {"name": r["지역"]}}, r["검색어"], dt)
                merged["근무요일/날짜"] = r["근무요일/날짜"] or merged["근무요일/날짜"]
                merged["근무시간"] = r["근무시간"] or merged["근무시간"]
                merged["급여"] = r["급여"] or merged["급여"]
                r.update(merged)
            if progress and i % 5 == 0:
                progress(i, len(rows))
    for r in rows:
        r["검색어"] = ", ".join(kw_of.get(r.get("_id"), [])) or r.get("검색어", "")
    if include_words:
        rows = [r for r in rows if _contains_any(r["제목"] + " " + r.get("내용", "") + " " + r["회사명"], include_words)]
    if exclude_words:
        rows = [r for r in rows if not _contains_any(r["제목"] + " " + r.get("내용", "") + " " + r["회사명"], exclude_words)]
    rows.sort(key=lambda x: x.get("_created") or datetime.min, reverse=True)
    return rows


def save_excel(rows: List[Dict[str, Any]], path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    widths = {"번호": 14, "제목": 44, "회사명": 18, "지역": 10, "주소": 30, "급여": 16, "고용형태": 8, "근무기간": 14,
              "업무": 16, "근무요일/날짜": 22, "근무시간": 14, "내용": 60, "전화연락": 8, "조회수": 7, "관심": 6,
              "지원자": 6, "등록일시": 16, "마감": 6, "검색어": 12, "당근링크": 40}
    wb = Workbook()
    ws = wb.active
    ws.title = "당근알바"
    ws.append(EXCEL_COLUMNS)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="FFE8D6")
        c.font = Font(bold=True, color="1F2937")
        c.alignment = Alignment(horizontal="center", vertical="center")
    link_font = Font(color="1D4ED8", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    for r in rows:
        ws.append([r.get(k, "") for k in EXCEL_COLUMNS])
        i = ws.max_row
        for j, k in enumerate(EXCEL_COLUMNS, 1):
            cell = ws.cell(row=i, column=j)
            if k == "당근링크" and cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = link_font
            cell.alignment = wrap if k in ("내용", "제목", "주소", "근무요일/날짜") else top
        if len(r.get("내용") or "") > 200:
            ws.row_dimensions[i].height = 90
    for j, k in enumerate(EXCEL_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(k, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return path
