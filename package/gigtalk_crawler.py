# -*- coding: utf-8 -*-
"""
긱톡(gigtalk.co.kr) 게시글 수집 - 데이터 모듈.

사실(2026-10-12 확인)
- 목록: GET /api/v2/app/posts  (로그인 불필요, JSON)
    파라미터: type(job/knowHow/site/daily/usedTrade), meetId, categoryIds(쉼표), searchItem, offset, size(최대 200 확인)
    정렬은 최신순 고정(orderBy/sort 는 값 검증에 걸려 못 씀). count = 조건 전체 건수.
- 커뮤니티: GET /api/v1/app/meets, 카테고리: GET /api/v1/app/categories?meetId=&postType=
- 상세: GET /api/v1/app/posts/{id}  (목록 행과 거의 같음 - 본문(memo/contents)도 목록에 이미 들어 있음)
- 웹 주소: https://gigtalk.co.kr/post/{id}
"""
from __future__ import annotations

import html as _html
import json
import os
import random
import re
import time
from datetime import date, datetime, timedelta
from typing import Any, Callable, Dict, Iterable, List, Optional, Set, Tuple

import requests

BASE = "https://gigtalk.co.kr"
API_POSTS = BASE + "/api/v2/app/posts"
API_POST = BASE + "/api/v1/app/posts/{id}"
API_MEETS = BASE + "/api/v1/app/meets"
API_CATEGORIES = BASE + "/api/v1/app/categories"
POST_URL = BASE + "/post/{id}"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "ko-KR,ko;q=0.9",
    "Referer": BASE + "/",
}

# (화면 이름, API type)
POST_TYPES: List[Tuple[str, str]] = [
    ("긱잡(일감·알바)", "job"),
    ("노하우", "knowHow"),
    ("서비스", "site"),
    ("일상·후기", "daily"),
    ("중고거래", "usedTrade"),
]
TYPE_KO: Dict[str, str] = {
    "job": "긱잡", "knowHow": "노하우", "site": "서비스", "daily": "일상·후기",
    "usedTrade": "중고거래", "groupBuying": "공동구매", "centerTrade": "센터거래",
}
TYPE_BY_LABEL: Dict[str, str] = {lbl: t for lbl, t in POST_TYPES}

# 사이트가 안 될 때도 고를 수 있게 넣어 둔 커뮤니티 목록 (2026-10-12 기준)
BUILTIN_MEETS: List[Dict[str, Any]] = [
    {"id": 18, "name": "재택모음"}, {"id": 3, "name": "온라인 셀러"}, {"id": 2, "name": "리뷰·체험단"},
    {"id": 1, "name": "어필리에이트"}, {"id": 4, "name": "앱테크"}, {"id": 5, "name": "재능마켓"},
    {"id": 7, "name": "교육·과외"}, {"id": 6, "name": "IT프리랜서"}, {"id": 9, "name": "방송"},
    {"id": 8, "name": "미용"}, {"id": 13, "name": "필라테스"}, {"id": 15, "name": "헬스"},
    {"id": 16, "name": "마사지·스파"}, {"id": 14, "name": "요가"}, {"id": 17, "name": "물리치료"},
    {"id": 10, "name": "음식·배달"}, {"id": 12, "name": "퀵·택배"},
]

PAYMENT_KO = {
    "hourly": "시급", "daily": "일급", "weekly": "주급", "monthly": "월급",
    "caseBy": "건별", "discussion": "협의", "yearly": "연봉", "etc": "", "none": "",
}
RECRUIT_KO = {"process": "모집중", "end": "마감", "close": "마감", "done": "마감"}
PLACE_KO = {"none": "", "home": "재택", "office": "출근", "both": "재택/출근", "remote": "재택"}

EXCEL_COLUMNS: List[str] = [
    "번호", "종류", "커뮤니티", "카테고리", "제목", "회사명/작성자", "지역", "급여/가격", "근무시간",
    "모집마감", "연락처", "출처", "원문링크", "내용", "등록일시", "조회수", "좋아요", "댓글", "긱톡링크",
]

PAGE_SIZE = 100


# ──────────────────────────────── 유틸 ────────────────────────────────
def strip_html(s: Optional[str]) -> str:
    """에디터 HTML → 읽기 좋은 평문 (목록 항목은 '- ', 문단/줄바꿈은 개행)."""
    if not s:
        return ""
    t = str(s)
    t = re.sub(r"(?is)<\s*(script|style)[^>]*>.*?</\s*\1\s*>", "", t)
    t = re.sub(r"(?i)<\s*li[^>]*>", "- ", t)
    t = re.sub(r"(?i)</\s*(li|p|div|h[1-6]|tr|ul|ol)\s*>", "\n", t)
    t = re.sub(r"(?i)<\s*br\s*/?>", "\n", t)
    t = re.sub(r"<[^>]+>", "", t)
    t = _html.unescape(t).replace("\xa0", " ")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n\s*\n+", "\n", t)
    return t.strip()


def to_kst(row: Dict[str, Any]) -> Optional[datetime]:
    """createdAt(UTC, '2026-10-12T09:30:54.000000Z') → 한국시간. orderCreatedAt(KST 문자열)이 있으면 그걸 우선."""
    s = row.get("orderCreatedAt")
    if s:
        try:
            return datetime.strptime(str(s)[:19], "%Y-%m-%d %H:%M:%S")
        except Exception:
            pass
    s = row.get("createdAt")
    if not s:
        return None
    try:
        base = datetime.strptime(str(s)[:19], "%Y-%m-%dT%H:%M:%S")
        return base + timedelta(hours=9)
    except Exception:
        return None


def _fmt_amount(v: Any) -> str:
    try:
        n = int(float(v))
        return f"{n:,}원"
    except Exception:
        return str(v or "")


def payment_text(ji: Dict[str, Any]) -> str:
    if not ji:
        return ""
    pt = ji.get("paymentType") or ""
    etc = (ji.get("paymentAmountEtc") or "").strip()
    a1, a2 = ji.get("paymentAmount1"), ji.get("paymentAmount2")
    a0 = ji.get("paymentAmount")
    label = PAYMENT_KO.get(pt, pt)
    if etc:
        return etc
    if pt == "discussion":
        return "협의"
    if a1 is not None or a2 is not None:
        if a1 is not None and a2 is not None and a1 != a2:
            return f"{label} {_fmt_amount(a1)[:-1]}~{_fmt_amount(a2)}".strip()
        return f"{label} {_fmt_amount(a1 if a1 is not None else a2)}".strip()
    if a0 is not None:
        return f"{label} {_fmt_amount(a0)}".strip()
    return label


def working_time_text(ji: Dict[str, Any]) -> str:
    if not ji:
        return ""
    wt = (ji.get("workingTime") or "").strip()
    days = [k for k, v in (("월", ji.get("isMon")), ("화", ji.get("isTue")), ("수", ji.get("isWed")),
                           ("목", ji.get("isThu")), ("금", ji.get("isFri")), ("토", ji.get("isSat")),
                           ("일", ji.get("isSun"))) if v]
    hours = ""
    if ji.get("startHour") or ji.get("endHour"):
        hours = f"{ji.get('startHour') or ''}~{ji.get('endHour') or ''}"
    parts = [p for p in (wt, ",".join(days), hours) if p]
    if ji.get("isWorkingTimeNego"):
        parts.append("시간협의")
    return " ".join(parts)


def region_text(row: Dict[str, Any]) -> str:
    out = []
    for a in row.get("addresses") or []:
        s = " ".join(x for x in [(a.get("sido") or "").strip(), (a.get("sigungu") or "").strip()] if x)
        extra = (a.get("roadAddress") or a.get("jibunAddress") or "").strip()
        if extra and extra not in s:
            s = (s + " " + extra).strip()
        if a.get("isPossibleToWorkFromHome"):
            s = (s + " (재택가능)").strip()
        if s and s not in out:
            out.append(s)
    return ", ".join(out)


def normalize(row: Dict[str, Any]) -> Dict[str, Any]:
    """API 행 → 엑셀 한 줄."""
    t = row.get("type") or ""
    ji = row.get("jobInfo") or {}
    ki = row.get("knowHowInfo") or {}
    di = row.get("dailyInfo") or {}
    ui = row.get("usedTradeInfo") or {}
    meet = (row.get("meet") or {}).get("name") or ""
    cat = (row.get("category") or {}).get("name") or ""
    prof = row.get("profile") or {}
    created = to_kst(row)

    title = ""
    author = ""
    pay = ""
    wtime = ""
    deadline = ""
    contact = row.get("phone") or ""
    ref = ""
    ref_url = ""
    body = ""
    if t == "job":
        title = ji.get("title") or row.get("search") or ""
        author = ji.get("companyName") or prof.get("name") or ""
        pay = payment_text(ji)
        wtime = working_time_text(ji)
        deadline = ji.get("recruitEndDate") or ""
        st = RECRUIT_KO.get(ji.get("recruitState") or "", "")
        if st and st != "모집중":
            deadline = (deadline + " " + st).strip()
        contact = ji.get("contact") or contact
        ref = ji.get("reference") or ""
        ref_url = ji.get("referenceUrl") or ""
        body = strip_html(ji.get("memo"))
        place = PLACE_KO.get(ji.get("placeType") or "", "")
        if ji.get("place"):
            place = (place + " " + str(ji["place"])).strip()
        if place:
            wtime = (wtime + " / " + place).strip(" /")
    elif t == "knowHow":
        title = ki.get("title") or row.get("search") or ""
        author = ki.get("author") or prof.get("name") or ""
        pay = payment_text(ki) if ki.get("paymentType") not in (None, "none") else ""
        ref = ki.get("reference") or ""
        ref_url = ki.get("referenceUrl") or ""
        body = strip_html(ki.get("memo"))
    elif t == "usedTrade":
        title = ui.get("title") or row.get("search") or ""
        author = prof.get("name") or ""
        if ui.get("price") is not None:
            pay = _fmt_amount(ui.get("price"))
            if ui.get("regularPrice") not in (None, ui.get("price")):
                pay += f" (정가 {_fmt_amount(ui.get('regularPrice'))})"
        contact = ui.get("contact") or contact
        ref = ui.get("reference") or ""
        ref_url = ui.get("referenceUrl") or ""
        body = strip_html(ui.get("memo"))
        deadline = {"process": "판매중", "end": "판매완료"}.get(ui.get("tradeState") or "", "")
    else:  # daily / site / 기타
        title = di.get("title") or (row.get("search") or "")[:80]
        author = prof.get("name") or ""
        ref_url = di.get("referenceUrl") or ""
        body = strip_html(di.get("contents"))
    if not title:
        title = (row.get("search") or "")[:80]
    if row.get("email"):
        contact = (contact + " " + row["email"]).strip()

    pid = row.get("id")
    return {
        "번호": pid,
        "종류": TYPE_KO.get(t, t),
        "커뮤니티": meet,
        "카테고리": cat,
        "제목": title.strip(),
        "회사명/작성자": author,
        "지역": region_text(row),
        "급여/가격": pay,
        "근무시간": wtime,
        "모집마감": deadline,
        "연락처": contact,
        "출처": ref,
        "원문링크": ref_url,
        "내용": body,
        "등록일시": created.strftime("%Y-%m-%d %H:%M") if created else "",
        "조회수": row.get("readCount") or 0,
        "좋아요": row.get("likeCount") or 0,
        "댓글": row.get("countComment") or 0,
        "긱톡링크": POST_URL.format(id=pid),
        "_created": created,
        "_type": t,
    }


# ──────────────────────────────── 클라이언트 ────────────────────────────────
class GigtalkClient:
    def __init__(self, log: Optional[Callable[[str], None]] = None, delay: Tuple[float, float] = (0.3, 0.9)):
        self.log = log or (lambda s: None)
        self.delay = delay
        self.s = requests.Session()
        self.s.headers.update(HEADERS)
        self._n = 0

    def _sleep(self) -> None:
        if self._n and self.delay[1] > 0:
            time.sleep(random.uniform(*self.delay))
        self._n += 1

    def _get(self, url: str, params: Optional[Dict[str, Any]] = None, tries: int = 3) -> Any:
        last = None
        for i in range(tries):
            try:
                self._sleep()
                r = self.s.get(url, params=params, timeout=30)
                if r.status_code == 429 or r.status_code >= 500:
                    raise RuntimeError(f"HTTP {r.status_code}")
                j = r.json()
                code = str(j.get("code", ""))
                if code and code != "0000":
                    msg = j.get("message") or code
                    if code == "0404":
                        return None
                    raise RuntimeError(f"긱톡 응답 오류: {msg}")
                if "errors" in j and not j.get("data"):
                    raise RuntimeError("긱톡 요청 값 오류: " + json.dumps(j.get("errors"), ensure_ascii=False)[:200])
                return j.get("data")
            except Exception as e:  # noqa: BLE001
                last = e
                if i < tries - 1:
                    wait = 3 * (i + 1)
                    self.log(f"  재시도 {i + 1}/{tries - 1} ({e}) {wait}초 대기")
                    time.sleep(wait)
        raise RuntimeError(f"긱톡 접속 실패: {last}")

    def meets(self) -> List[Dict[str, Any]]:
        d = self._get(API_MEETS) or {}
        rows = d.get("rows") if isinstance(d, dict) else d
        out = [{"id": r["id"], "name": r.get("name") or str(r["id"])} for r in (rows or []) if r.get("state", "active") == "active"]
        return out or list(BUILTIN_MEETS)

    def categories(self, meet_id: int, post_type: Optional[str] = None) -> List[Dict[str, Any]]:
        p: Dict[str, Any] = {"meetId": meet_id}
        if post_type:
            p["postType"] = post_type
        d = self._get(API_CATEGORIES, p)
        if not d:
            return []
        rows = d.get("rows") if isinstance(d, dict) else d
        return [{"id": r["id"], "name": r.get("name") or "", "postType": r.get("postType") or "", "meetId": meet_id}
                for r in (rows or []) if r.get("state", "active") == "active"]

    def posts(self, post_type: str, meet_id: Optional[int] = None, category_ids: Optional[Iterable[int]] = None,
              search: Optional[str] = None, offset: int = 0, size: int = PAGE_SIZE) -> Tuple[int, List[Dict[str, Any]]]:
        p: Dict[str, Any] = {"type": post_type, "offset": offset, "size": size}
        if meet_id:
            p["meetId"] = meet_id
        cids = [int(c) for c in (category_ids or [])]
        if cids:
            p["categoryIds"] = ",".join(str(c) for c in cids)
        if search:
            p["searchItem"] = search
        d = self._get(API_POSTS, p) or {}
        return int(d.get("count") or 0), list(d.get("rows") or [])

    def post_detail(self, post_id: int) -> Dict[str, Any]:
        return self._get(API_POST.format(id=post_id)) or {}


# ──────────────────────────────── 수집 ────────────────────────────────
def _contains_any(text: str, words: List[str]) -> bool:
    t = (text or "").lower()
    return any(w.lower() in t for w in words if w)


def collect(
    client: GigtalkClient,
    post_types: List[str],
    meet_ids: List[int],
    category_ids: List[int],
    search_terms: List[str],
    start: date,
    end: date,
    max_count: int = 500,
    include_words: Optional[List[str]] = None,
    exclude_words: Optional[List[str]] = None,
    region_words: Optional[List[str]] = None,
    only_open: bool = False,
    seen_ids: Optional[Set[int]] = None,
    stop: Optional[Callable[[], bool]] = None,
    log: Optional[Callable[[str], None]] = None,
    progress: Optional[Callable[[int, int], None]] = None,
    category_meet: Optional[Dict[int, int]] = None,
) -> List[Dict[str, Any]]:
    """조건에 맞는 글을 최신순으로 모아 엑셀 행(normalize 결과)으로 돌려준다.

    - 목록은 최신순 고정이라 등록일이 start 보다 오래된 행이 연속으로 나오면 그 조합은 중단.
    - category_meet: 카테고리 id → 커뮤니티 id (카테고리를 골랐을 때 커뮤니티별로 묶어서 요청).
    """
    log = log or (lambda s: None)
    stop = stop or (lambda: False)
    include_words = [w for w in (include_words or []) if w]
    exclude_words = [w for w in (exclude_words or []) if w]
    region_words = [w for w in (region_words or []) if w]
    seen_ids = seen_ids or set()
    terms: List[Optional[str]] = [t for t in (search_terms or []) if t] or [None]
    today = date.today()

    # 요청 조합 만들기: (type, meetId, categoryIds)
    combos: List[Tuple[str, Optional[int], List[int]]] = []
    for t in post_types:
        if category_ids:
            by_meet: Dict[Optional[int], List[int]] = {}
            for c in category_ids:
                m = (category_meet or {}).get(c)
                by_meet.setdefault(m, []).append(c)
            for m, cs in by_meet.items():
                combos.append((t, m, cs))
        elif meet_ids:
            for m in meet_ids:
                combos.append((t, m, []))
        else:
            combos.append((t, None, []))

    got: Dict[int, Dict[str, Any]] = {}
    order: List[int] = []
    total_est = 0
    for (t, m, cs) in combos:
        for term in terms:
            if stop():
                break
            offset = 0
            older_streak = 0
            label = f"{TYPE_KO.get(t, t)}" + (f" / 커뮤니티 {m}" if m else "") + (f" / 카테고리 {len(cs)}개" if cs else "") + (f" / 검색 '{term}'" if term else "")
            count = None
            while not stop():
                try:
                    count, rows = client.posts(t, m, cs, term, offset=offset, size=PAGE_SIZE)
                except Exception as e:  # noqa: BLE001
                    log(f"  [{label}] 목록 실패: {e}")
                    break
                if offset == 0:
                    log(f"[{label}] 전체 {count}건, 최신순으로 확인")
                    total_est += count
                if not rows:
                    break
                for r in rows:
                    created = to_kst(r)
                    d = created.date() if created else None
                    if d and d < start:
                        older_streak += 1
                        if older_streak >= 30:
                            break
                        continue
                    older_streak = 0
                    if d and d > end:
                        continue
                    pid = r.get("id")
                    if pid in got or pid in seen_ids:
                        continue
                    if r.get("isHidden") or r.get("deletedAt"):
                        continue
                    n = normalize(r)
                    hay = " ".join(str(n.get(k) or "") for k in ("제목", "내용", "회사명/작성자", "카테고리"))
                    if include_words and not _contains_any(hay, include_words):
                        continue
                    if exclude_words and _contains_any(hay, exclude_words):
                        continue
                    if region_words and not _contains_any(n.get("지역") or "", region_words):
                        continue
                    if only_open and t == "job":
                        dl = (n.get("모집마감") or "").strip()
                        if "마감" in dl:
                            continue
                        try:
                            if dl and date.fromisoformat(dl[:10]) < today:
                                continue
                        except Exception:
                            pass
                    got[pid] = n
                    order.append(pid)
                    if max_count and len(order) >= max_count:
                        break
                if progress:
                    progress(len(order), max_count or 0)
                if older_streak >= 30 or (max_count and len(order) >= max_count):
                    break
                offset += len(rows)
                if count is not None and offset >= count:
                    break
                if len(rows) < PAGE_SIZE:
                    break
            log(f"  [{label}] 누적 {len(order)}건")
            if max_count and len(order) >= max_count:
                break
        if max_count and len(order) >= max_count:
            log(f"최대 건수 {max_count}건에 도달해 멈춥니다.")
            break
    out = [got[i] for i in order]
    out.sort(key=lambda x: x.get("_created") or datetime.min, reverse=True)
    return out


# ──────────────────────────────── 저장 ────────────────────────────────
def load_seen(path: str) -> Dict[str, str]:
    try:
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as f:
                d = json.load(f) or {}
                return {str(k): str(v) for k, v in d.items()}
    except Exception:
        pass
    return {}


def save_seen(path: str, seen: Dict[str, str], keep_days: int = 180) -> None:
    cutoff = (date.today() - timedelta(days=keep_days)).isoformat()
    trimmed = {k: v for k, v in seen.items() if (v or "9999") >= cutoff}
    try:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(trimmed, f, ensure_ascii=False)
    except Exception:
        pass


def save_excel(rows: List[Dict[str, Any]], path: str, split_by_type: bool = True) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    widths = {
        "번호": 9, "종류": 8, "커뮤니티": 11, "카테고리": 14, "제목": 48, "회사명/작성자": 20, "지역": 18,
        "급여/가격": 16, "근무시간": 20, "모집마감": 11, "연락처": 16, "출처": 10, "원문링크": 30, "내용": 60,
        "등록일시": 16, "조회수": 7, "좋아요": 7, "댓글": 6, "긱톡링크": 32,
    }
    head_fill = PatternFill("solid", fgColor="E8EEF9")
    head_font = Font(bold=True, color="1F2937")
    link_font = Font(color="1D4ED8", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    def fill(ws, items: List[Dict[str, Any]]) -> None:
        ws.append(EXCEL_COLUMNS)
        for c in ws[1]:
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="center", vertical="center")
        for r in items:
            ws.append([r.get(k, "") for k in EXCEL_COLUMNS])
            i = ws.max_row
            for j, k in enumerate(EXCEL_COLUMNS, 1):
                cell = ws.cell(row=i, column=j)
                if k in ("원문링크", "긱톡링크") and cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.font = link_font
                cell.alignment = wrap if k in ("내용", "제목", "근무시간", "지역") else top
            if len(r.get("내용") or "") > 200:
                ws.row_dimensions[i].height = 90
        for j, k in enumerate(EXCEL_COLUMNS, 1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(k, 12)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    ws = wb.active
    ws.title = "전체"
    fill(ws, rows)
    if split_by_type:
        kinds: Dict[str, List[Dict[str, Any]]] = {}
        for r in rows:
            kinds.setdefault(r.get("종류") or "기타", []).append(r)
        if len(kinds) > 1:
            for k, items in sorted(kinds.items(), key=lambda kv: -len(kv[1])):
                fill(wb.create_sheet(title=re.sub(r"[\\/*?:\[\]]", "_", k)[:28]), items)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return path
