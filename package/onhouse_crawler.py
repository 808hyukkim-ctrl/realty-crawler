"""
OnHouse 부동산 사이트 크롤러
"""
import logging
import re
import json
from curl_cffi import requests as curl_req
from bs4 import BeautifulSoup
from typing import Optional, Dict, Any, List, Union, Callable


class OnhouseCrawler:
    LOGIN_URL = "https://www.onhouse.com/index.php/dataFunction/login"
    SEARCH_URL = "https://www.onhouse.com/index.php/dataFunction/rentMapList"
    SEARCH_PAGE_URL = "https://www.onhouse.com/index.php/dataFunction/rentMapList/page"
    DETAIL_URL = "https://www.onhouse.com/index/rent_view"

    def __init__(self, id: str = "", pwd: str = "", verbose: bool = True):
        self.host = "https://www.onhouse.com"
        self.session = curl_req.Session()
        self.timeout_sec = 20
        self.session.headers.update({
            "accept": "*/*",
            "accept-encoding": "gzip, deflate, br, zstd",
            "accept-language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
            "content-type": "application/x-www-form-urlencoded; charset=UTF-8",
            "origin": self.host,
            "priority": "u=1, i",
            "referer": f"{self.host}/index/sale_map",
            "sec-ch-ua": '"Not(A:Brand";v="8", "Chromium";v="144", "Google Chrome";v="144"',
            "sec-ch-ua-mobile": "?0",
            "sec-ch-ua-platform": '"Windows"',
            "sec-fetch-dest": "empty",
            "sec-fetch-mode": "cors",
            "sec-fetch-site": "same-origin",
            "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/144.0.0.0 Safari/537.36",
            "x-requested-with": "XMLHttpRequest",
        })
        self.id = id
        self.pwd = pwd
        self._logged_in = False
        self.verbose = verbose

    def _log(self, msg: str):
        if not self.verbose:
            return
        print(f"[OnhouseCrawler] {msg}", flush=True)

    def login(
        self,
        id: Optional[str] = None,
        pwd: Optional[str] = None
    ) -> Dict[str, Any]:
        
        login_id = id or self.id
        payload = {
            "from_page": "",
            "dg_login_ver": "",
            "id": login_id,
            "pwd": pwd or self.pwd,
        }
        self._log(f"login start id={login_id}")
        r = self.session.post(
            self.LOGIN_URL,
            data=payload,
            impersonate="chrome",
            timeout=self.timeout_sec,
        )
        r.raise_for_status()
        # 로그인 후 세션 설정용 GET 요청
        self.session.get(
            f"{self.host}/login_opt/toLoginNormal",
            impersonate="chrome",
            timeout=self.timeout_sec,
        )
        self._logged_in = True
        self._log(f"login success status={r.status_code}")
        try:
            return r.json()
        except Exception:
            return {"text": r.text, "status_code": r.status_code}

    @property
    def is_logged_in(self) -> bool:
        return self._logged_in

    def _parse_ids_from_html(self, html: str) -> List[str]:
        """list_body HTML에서 rent_view URL용 ID 추출"""
        ids = set()
        # id="listItem_3328524"
        for m in re.finditer(r'listItem_(\d+)', html):
            ids.add(m.group(1))
        # rent_view/3328524
        for m in re.finditer(r'rent_view/(\d+)', html):
            ids.add(m.group(1))
        # openListItemInfo('3328524'
        for m in re.finditer(r"openListItemInfo\s*\(\s*['\"]?(\d+)", html):
            ids.add(m.group(1))
        # chkItem value="3328524"
        for m in re.finditer(r'chkItem[^>]*value=["\'](\d+)', html):
            ids.add(m.group(1))
        return sorted(ids, key=int)

    def search(
        self,
        sw_lat: float = 37.3899004,
        ne_lat: float = 37.3900004,
        sw_lng: float = 127.0985522,
        ne_lng: float = 127.0986522,
        trade_type: Union[str, List[str]] = "'매매'",
        room_type: str = "all",
        limit: int = 30,
        level: int = 7,
        page_step: int = 30,
        page: int = 0,
        fetch_details: bool = True,
        refer_min_price: Optional[int] = None,
        refer_max_price: Optional[int] = None,
        min_price: Optional[int] = None,
        max_price: Optional[int] = None,
        min_month_price: Optional[int] = None,
        max_month_price: Optional[int] = None,
        min_area: Optional[float] = None,
        max_area: Optional[float] = None,
        is_cancelled: Optional[Callable[[], bool]] = None,
        **kwargs,
    ) -> List[str] | List[Dict[str, Any]]:
        """
        rentMapList/page/{offset} URL로 지정 페이지 1페이지만 조회. (페이지 순회는 호출부에서)
        refer*: 기준가(만원), min/maxPrice: 보증금/매매가, *MonthPrice: 월세, *Area: 평수

        Args:
            trade_type: "'매매'" 또는 ["월세","매매"] → "'월세','매매'"
            room_type: 전체(all), 주택, 오피, 주택/오피, 사무실, 상가, 사무실/상가, 분양사무실
            page: 조회할 페이지 (0부터)
            fetch_details: True면 해당 페이지 ID 수집 후 즉시 detail 크롤링, List[Dict] 반환
        """
        if isinstance(trade_type, list):
            trade_type_str = ",".join(f"'{t}'" for t in trade_type) if trade_type else "all"
        else:
            trade_type_str = trade_type

        # roomType: all이면 그대로, 한글이면 '주택' 형식으로
        room_type_str = room_type if (room_type or "").lower() == "all" else f"'{room_type}'"

        payload = {
            "device": "pc",
            "showType": "확인일",
            "text": "",
            "buildingIdx": "",
            "dongCode": "",
            "bunjiCode": "",
            "roomType": room_type_str,
            "saleType": "all",
            "structure": "all",
            "tradeType": trade_type_str,
            "shortTermYn": "",
            "referMinPrice": str(refer_min_price) if refer_min_price is not None else "",
            "referMaxPrice": str(refer_max_price) if refer_max_price is not None else "",
            "minPrice": str(min_price) if min_price is not None else "",
            "maxPrice": str(max_price) if max_price is not None else "",
            "minMonthPrice": str(min_month_price) if min_month_price is not None else "",
            "maxMonthPrice": str(max_month_price) if max_month_price is not None else "",
            "managementFeeYn": "",
            "minArea": str(min_area) if min_area is not None else "",
            "maxArea": str(max_area) if max_area is not None else "",
            "minFloor": "",
            "maxFloor": "",
            "floorArray": "",
            "rooftop": "",
            "compYear": "all",
            "img": "all",
            "noDeposit": "",
            "fulloption": "",
            "noSemiBasement": "",
            "noFirstFloor": "",
            "elevator": "",
            "interior": "",
            "noPremiumPrice": "",
            "pet": "",
            "parking": "",
            "roomCnt": "all",
            "cfLoan": "",
            "moveInType": "",
            "level": str(level),
            "swLat": sw_lat,
            "neLat": ne_lat,
            "swLng": sw_lng,
            "neLng": ne_lng,
            "limit": limit,
            "order1": "R.CHK_DATE|DESC",
            "order2": "",
            "status": "거래가능",
            "phone": "",
            "jibun_addr": "",
            "building_name": "",
            "zerooption": "",
            "view_type": "hori",
        }
        payload.update(kwargs)
        if is_cancelled and is_cancelled():
            return []

        offset = page * page_step
        url = f"{self.SEARCH_PAGE_URL}/{offset if offset else ''}"
        # self._log(
        #     "search request "
        #     f"page={page} offset={offset} limit={limit} tradeType={trade_type_str} roomType={room_type_str} "
        #     f"refer=({payload['referMinPrice']},{payload['referMaxPrice']}) "
        #     f"price=({payload['minPrice']},{payload['maxPrice']}) "
        #     f"month=({payload['minMonthPrice']},{payload['maxMonthPrice']}) "
        #     f"area=({payload['minArea']},{payload['maxArea']})"
        # )
        r = self.session.post(url, data=payload, impersonate="chrome", timeout=self.timeout_sec)
        r.raise_for_status()
        text = r.text
        # self._log(f"search response status={r.status_code} body_len={len(text)}")
        # self._log(payload); import sys; sys.exit(1)
        if text.strip().startswith("<") or "list_body" in text or "listItem_" in text:
            page_ids = self._parse_ids_from_html(text)
        else:
            try:
                data = r.json()
                if isinstance(data, list):
                    page_ids = [str(x) for x in data]
                elif isinstance(data, dict) and "ids" in data:
                    page_ids = [str(x) for x in data["ids"]]
                else:
                    page_ids = []
            except Exception:
                page_ids = []
        self._log(f"search parsed_ids={len(page_ids)} page={page}")

        # print(page_ids)
        if fetch_details:
            self._log(f"search fetch_details start count={len(page_ids)}")
            all_rows: List[Dict[str, Any]] = []
            for i, hid in enumerate(page_ids, 1):
                if is_cancelled and is_cancelled():
                    break
                try:
                    self._log(f"  detail[{i}/{len(page_ids)}] id={hid} start")
                    row = self.crawl_detail(str(hid), is_cancelled=is_cancelled)
                    all_rows.append(row)
                    self._log(f"  detail[{i}/{len(page_ids)}] id={hid} ok keys={len(row.keys())}")
                except Exception as e:
                    print(f"  detail {hid} 실패: {e}")
                    all_rows.append({"매물ID": hid, "URL": f"{self.DETAIL_URL}/{hid}", "오류": str(e)})
            self._log(f"search fetch_details done rows={len(all_rows)}")
            return all_rows
        return sorted(page_ids, key=int)

    def search_all_pages_with_details(
        self,
        sw_lat: float = 37.3899004,
        ne_lat: float = 37.3900004,
        sw_lng: float = 127.0985522,
        ne_lng: float = 127.0986522,
        trade_type: Union[str, List[str]] = "'매매'",
        room_type: str = "all",
        limit: int = 30,
        page_step: int = 30,
        refer_min_price: Optional[int] = None,
        refer_max_price: Optional[int] = None,
        min_price: Optional[int] = None,
        max_price: Optional[int] = None,
        min_month_price: Optional[int] = None,
        max_month_price: Optional[int] = None,
        min_area: Optional[float] = None,
        max_area: Optional[float] = None,
        progress_callback: Optional[Any] = None,
        **kwargs,
    ) -> List[Dict[str, Any]]:
        """전체 페이지 순회하며 각 페이지마다 detail 크롤링 후 반환"""
        self._log("search_all_pages_with_details start")
        all_rows: List[Dict[str, Any]] = []
        page_num = 0
        while True:
            self._log(f"page loop start page={page_num}")
            result = self.search(
                sw_lat=sw_lat,
                ne_lat=ne_lat,
                sw_lng=sw_lng,
                ne_lng=ne_lng,
                trade_type=trade_type,
                room_type=room_type,
                limit=limit,
                page_step=page_step,
                page=page_num,
                fetch_details=True,
                refer_min_price=refer_min_price,
                refer_max_price=refer_max_price,
                min_price=min_price,
                max_price=max_price,
                min_month_price=min_month_price,
                max_month_price=max_month_price,
                min_area=min_area,
                max_area=max_area,
                **kwargs,
            )
            if not result:
                self._log(f"page loop stop page={page_num} reason=empty")
                break
            all_rows.extend(result)
            self._log(f"page loop done page={page_num} rows={len(result)} total={len(all_rows)}")
            if progress_callback:
                progress_callback(page_num, len(result), len(all_rows))
            if len(result) < limit:
                self._log(f"page loop stop page={page_num} reason=rows_lt_limit")
                break
            page_num += 1
        self._log(f"search_all_pages_with_details finished total={len(all_rows)}")
        return all_rows

    def fetch_detail(self, id: str) -> str:
        """상세 페이지 HTML 가져오기"""
        url = f"{self.DETAIL_URL}/{id}"
        self._log(f"fetch_detail id={id} url={url}")
        r = self.session.get(url, impersonate="chrome", timeout=self.timeout_sec)
        r.raise_for_status()
        self._log(f"fetch_detail done id={id} status={r.status_code} body_len={len(r.text)}")
        return r.text

    def _parse_detail_html(self, html: str, id: str) -> Dict[str, Any]:
        """상세 페이지 HTML에서 데이터 추출"""
        soup = BeautifulSoup(html, "html.parser")
        out: Dict[str, Any] = {"매물ID": id, "URL": f"{self.DETAIL_URL}/{id}"}

        # 물건번호: topInfoArea .left_box
        top_info = soup.find("div", class_="topInfoArea")
        if top_info:
            left = top_info.find("div", class_="left_box")
            if left:
                txt = left.get_text(strip=True)
                m = re.search(r"No\s*(\d+)", txt)
                if m:
                    out["물건번호"] = f"No{m.group(1)}"

        # 건물명, 주소, 지번, 호실: topMainTitleArea
        title_area = soup.find("div", class_="topMainTitleArea")
        if title_area:
            left = title_area.find("div", class_="left_box")
            if left:
                bar = left.find("span", class_="title_bar")
                if bar:
                    prev = bar.previous_sibling
                    out["건물명"] = str(prev).strip() if prev else ""
                title_addr = left.find("span", class_="title_addr")
                if title_addr:
                    jibun_el = title_addr.find(id="jibun_hide")
                    if jibun_el and jibun_el.get("data-jibun"):
                        out["지번"] = jibun_el["data-jibun"]
                    out["주소_호실"] = title_addr.get_text(separator=" ", strip=True)
            addr_title = title_area.find("h6", class_="addr_title")
            if addr_title:
                out["전체주소"] = addr_title.get_text(strip=True)

        # bodyTop: bodyBox (거래유형/가격, 전용면적, 입주가능일, 해당층/전체층, 확인일/등록일)
        body_top = soup.find("div", class_="bodyTop")
        if body_top:
            boxes = body_top.find_all("div", class_="bodyBox")
            for box in boxes:
                # boxFlex (확인일, 등록일 등)
                flex = box.find("div", class_="boxFlex")
                if flex:
                    titles = flex.find_all("div", class_="flex_title")
                    descs = flex.find_all("div", class_="flex_desc")
                    for t, d in zip(titles, descs):
                        k = t.get_text(strip=True)
                        v = d.get_text(strip=True)
                        if k:
                            out[k] = v
                    continue
                title_el = box.find("div", class_="box_title")
                desc_el = box.find("div", class_="box_desc")
                if not title_el:
                    continue
                title = title_el.get_text(strip=True)
                if not title:
                    continue
                val = desc_el.get_text(strip=True) if desc_el else ""
                out[title] = val
                sub = box.find("div", class_="box_sub")
                if sub:
                    tags = []
                    for span in sub.find_all(["span"], class_=lambda c: c and ("border_tag" in c or "bg_tag" in c)):
                        tags.append(span.get_text(strip=True))
                    if tags:
                        out[f"{title}_태그"] = ", ".join(tags)

        # bodyBottom: 주차, 엘리베이터 등
        body_bottom = soup.find("div", class_="bodyBottom")
        if body_bottom:
            for box in body_bottom.find_all("div", class_="bodyBox"):
                title_el = box.find("div", class_="box_title")
                desc_el = box.find("div", class_="box_desc")
                if title_el and desc_el:
                    k = title_el.get_text(strip=True)
                    v = desc_el.get_text(strip=True)
                    if k:
                        out[k] = v

        return out

    def _sanitize_excel(self, s: str) -> str:
        """엑셀 제어문자 제거"""
        if not s:
            return ""
        return "".join(
            c if (ord(c) in (0x09, 0x0A, 0x0D) or 0x20 <= ord(c) <= 0x7E or ord(c) >= 0xA0) else " "
            for c in str(s)
        )

    def crawl_detail(self, id: str, is_cancelled: Optional[Callable[[], bool]] = None) -> Dict[str, Any]:
        """상세 페이지 크롤링 후 추출 데이터 반환"""
        if is_cancelled and is_cancelled():
            raise InterruptedError("중단 요청됨")
        self._log(f"crawl_detail start id={id}")
        html = self.fetch_detail(id)
        with open("html.html", "w", encoding="utf-8") as f:
            f.write(html)
        out = self._parse_detail_html(html, id)
        self._log(f"crawl_detail done id={id} keys={len(out.keys())}")
        return out

    def crawl_details_to_excel(
        self,
        ids: Optional[List[str]] = None,
        rows: Optional[List[Dict[str, Any]]] = None,
        output_path: str = "onhouse_detail.xlsx",
    ) -> str:
        """ID 리스트 순회하며 상세 페이지 크롤링 후 엑셀 저장. rows가 있으면 그대로 저장"""
        self._log(
            "crawl_details_to_excel start "
            f"ids={0 if ids is None else len(ids)} rows={0 if rows is None else len(rows)} "
            f"output={output_path}"
        )
        try:
            import openpyxl
        except ImportError:
            raise ImportError("pip install openpyxl")

        if rows is None:
            rows = []
            for i, hid in enumerate(ids or [], 1):
                print(f"[{i}/{len(ids)}] {hid}")
                try:
                    data = self.crawl_detail(str(hid))
                    rows.append(data)
                except Exception as e:
                    print(f"  실패: {e}")
                    rows.append({"매물ID": hid, "URL": f"{self.DETAIL_URL}/{hid}", "오류": str(e)})

        if not rows:
            raise ValueError("저장할 데이터 없음")

        all_keys = []
        seen = set()
        for r in rows:
            for k in r.keys():
                if k not in seen:
                    seen.add(k)
                    all_keys.append(k)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "매물상세"
        ws.append([self._sanitize_excel(str(k)) for k in all_keys])
        for r in rows:
            ws.append([self._sanitize_excel(str(r.get(k, ""))) for k in all_keys])
        self._apply_hyperlinks(ws, all_keys)
        wb.save(output_path)
        self._log(f"crawl_details_to_excel saved rows={len(rows)} path={output_path}")
        return output_path

    def _apply_hyperlinks(self, ws, all_keys: List[str]) -> None:
        """URL/링크 컬럼을 엑셀 하이퍼링크 셀로 저장"""
        try:
            from openpyxl.styles import Font
        except Exception:
            return
        url_col_indices = []
        for i, key in enumerate(all_keys, 1):
            k = str(key).strip().lower()
            if ("url" in k) or ("link" in k) or ("링크" in str(key)):
                url_col_indices.append(i)
        if not url_col_indices:
            return
        link_font = Font(color="0563C1", underline="single")
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in url_col_indices:
                cell = ws.cell(row=row_idx, column=col_idx)
                value = str(cell.value or "").strip()
                if value.lower().startswith(("http://", "https://")):
                    cell.hyperlink = value
                    cell.font = link_font


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    # 로컬 테스트용. 계정은 소스에 넣지 말고 환경변수로 준다: set ONHOUSE_ID=... / set ONHOUSE_PW=...
    import os as _os
    crawler = OnhouseCrawler(id=_os.environ.get("ONHOUSE_ID", ""), pwd=_os.environ.get("ONHOUSE_PW", ""))
    res = crawler.login()
    rows = crawler.search_all_pages_with_details(
        progress_callback=lambda p, c, t: print(f"  page {p}: {c}건 (누적 {t}건)"),
    )
    print(f"매물 {len(rows)}개")
    if rows:
        path = crawler.crawl_details_to_excel(rows=rows, output_path="onhouse_detail.xlsx")
        print(f"저장: {path}")