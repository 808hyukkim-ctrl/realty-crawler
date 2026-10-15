"""
app_v1_main / app_v1_site_main / app_v2_main 에서 공통으로 쓰는 유틸·위젯.
"""
from __future__ import annotations

import os
import re
import subprocess
import sys
from typing import Any, List, Optional

from PySide6.QtCore import QEvent, Qt, Signal
from PySide6.QtGui import (
    QColor,
    QFont,
    QFontMetrics,
    QIntValidator,
    QPainter,
    QPen,
    QStandardItem,
    QStandardItemModel,
)
from PySide6.QtWidgets import (
    QCheckBox,
    QComboBox,
    QFrame,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QVBoxLayout,
    QWidget,
)


def parse_keywords_csv(raw: str) -> List[str]:
    return [x.strip() for x in str(raw or "").split(",") if x.strip()]


def keyword_regex(keywords: List[str]):
    """키워드 OR 정규식. 영문 키워드는 앞뒤가 영문이면 매칭하지 않는다.

    LH/SH/HUG 처럼 짧은 영문 키워드가 English, flash, shop, hugging 같은 단어 속에
    섞여 잡히던 문제(2026-09-23) 때문에 영문 키워드에는 글자 경계를 준다.
    한글 키워드(대출, 보증보험 등)는 붙여 쓰는 경우가 많아 그대로 부분일치한다.
    """
    import re as _re

    kws = [str(k).strip() for k in (keywords or []) if str(k).strip()]
    if not kws:
        return None
    parts = []
    for k in sorted(kws, key=len, reverse=True):
        esc = _re.escape(k)
        if _re.fullmatch(r"[A-Za-z]+", k):
            esc = r"(?<![A-Za-z])" + esc + r"(?![A-Za-z])"
        parts.append(esc)
    return _re.compile("|".join(parts), _re.IGNORECASE)


def contains_any_keyword(text: Any, keywords: List[str]) -> bool:
    """키워드 중 하나라도 본문에 있으면 True (OR). 영문 대소문자는 무시한다 (lh == LH).
    영문 키워드는 단어 경계를 지켜 English 의 sh 같은 오탐을 걸러낸다."""
    rx = keyword_regex(keywords)
    if rx is None:
        return True
    return bool(rx.search(str(text or "")))


# ---------------------------------------------------------------------------
# 수집 결과 엑셀을 종류별 시트로 나누기 (2026-10-08, 사용자: "원룸은 원룸대로 투룸은 투룸대로 사무실·상가… 카테고리별로 엑셀 정리")
#   첫 시트는 '전체'(원본 그대로), 그 뒤에 종류별 시트를 건수 많은 순으로 붙인다. 서식·하이퍼링크·키워드 강조(리치텍스트)도 같이 복사.
CATEGORY_COLUMNS = ("종류", "매물유형", "건축물용도")          # 네이버 '종류', 당근 '매물유형'
CATEGORY_TAG_COLUMNS = ("전세_태그", "월세_태그", "매매_태그")   # 온하우스는 태그 첫 항목이 방 종류 ("투룸, 2002년")


def normalize_category(value) -> str:
    """'원룸(분리형)'·'오픈형원룸'·'주방분리형원룸' → '원룸', '쓰리룸+' → '쓰리룸', '단독/다가구' → '단독·다가구'. 빈 값은 '기타'."""
    v = re.sub(r"\(.*?\)", "", str(value if value is not None else "")).strip()
    v = v.split(",")[0].strip()
    if not v:
        return "기타"
    for key in ("원룸", "투룸", "쓰리룸", "사무실", "상가", "오피스텔", "아파트", "빌라"):
        if key in v:
            return key
    return re.sub(r"[\[\]:*?/\\]", "·", v)[:31]


def split_excel_by_category(path: str) -> List[str]:
    """저장된 xlsx 의 첫 시트를 종류별 시트로 나눠 같은 파일에 덧붙인다. 돌려주는 값: 만든 시트 이름들(나눌 게 없으면 빈 목록).
    실패해도 예외를 밖으로 내지 않는다 — 엑셀 저장 자체는 이미 끝난 뒤라 수집 결과에 영향이 없어야 한다."""
    try:
        import openpyxl
        from copy import copy

        wb = openpyxl.load_workbook(path, rich_text=True)
        ws = wb.worksheets[0]
        headers = [str(c.value or "").strip() for c in ws[1]]
        cat_cols = [headers.index(h) for h in CATEGORY_COLUMNS if h in headers]
        tag_cols = [headers.index(h) for h in CATEGORY_TAG_COLUMNS if h in headers]
        if not cat_cols and not tag_cols:
            return []

        def category_of(cells) -> Optional[str]:
            if not any(c.value not in (None, "") for c in cells):
                return None                                  # 빈 행
            if cat_cols:
                return normalize_category(cells[cat_cols[0]].value)
            for i in tag_cols:
                if cells[i].value not in (None, ""):
                    return normalize_category(cells[i].value)
            return "기타"

        groups: dict = {}
        for row in ws.iter_rows(min_row=2):
            cat = category_of(row)
            if cat is None:
                continue
            groups.setdefault(cat, []).append(row)
        if len(groups) <= 1:
            return []                                        # 한 종류뿐이면 나눌 필요 없음

        for extra in wb.worksheets[1:]:                       # 다시 돌릴 때 이전 종류 시트 제거
            wb.remove(extra)
        ws.title = "전체"
        made: List[str] = []
        used = {"전체"}
        for cat, rows in sorted(groups.items(), key=lambda kv: -len(kv[1])):
            title = cat
            n = 2
            while title in used:
                title = f"{cat}{n}"; n += 1
            used.add(title)
            new = wb.create_sheet(title)
            for col_letter, dim in ws.column_dimensions.items():
                if dim.width:
                    new.column_dimensions[col_letter].width = dim.width
            r_out = 0
            for src_row in [ws[1], *rows]:
                r_out += 1
                for cell in src_row:
                    dst = new.cell(row=r_out, column=cell.column)
                    dst.value = cell.value
                    if cell.has_style:
                        dst._style = copy(cell._style)
                    if cell.hyperlink:
                        dst.hyperlink = copy(cell.hyperlink)
            new.freeze_panes = "A2"
            made.append(title)
        wb.save(path)
        return made
    except Exception:
        return []


def parse_saved_output_path_from_finish_message(msg: str) -> Optional[str]:
    """Worker finished 메시지에서 저장된 파일 경로 추출. (저장 완료|중단 저장 완료): path (N건) 형식."""
    s = str(msg or "").strip()
    if "저장 완료" not in s and "중단 저장 완료" not in s:
        return None
    m = re.match(
        r"^(?:중단 저장 완료|저장 완료)\s*:\s*(.+)\s+\(\d+건\)\s*$",
        s,
    )
    if not m:
        return None
    path = os.path.normpath(m.group(1).strip().strip('"'))
    return path if os.path.isfile(path) else None


def open_path_in_os(path: str) -> None:
    """파일을 OS 기본 연결 프로그램으로 연다."""
    p = os.path.abspath(os.path.expanduser(str(path)))
    if not os.path.isfile(p):
        raise FileNotFoundError(p)
    if sys.platform == "win32":
        os.startfile(p)  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        subprocess.run(["open", p], check=False)
    else:
        subprocess.run(["xdg-open", p], check=False)


DATE_PRESETS = ["전체", "오늘", "어제/오늘", "일주일", "한달"]
DATE_PRESET_DAYS = {"오늘": 0, "어제/오늘": 1, "일주일": 7, "한달": 30}


def date_preset_to_range(preset: Any) -> tuple[Optional[str], Optional[str]]:
    """'오늘'/'어제/오늘'/'일주일'/'한달' → ("yyyyMMdd", "yyyyMMdd"), '전체'·미지정 → (None, None).
    끝은 항상 오늘, 시작은 오늘에서 N일 전(네이버 탭의 등록일 프리셋과 같은 기준)."""
    from datetime import date, timedelta

    days = DATE_PRESET_DAYS.get(str(preset or "").strip())
    if days is None:
        return None, None
    today = date.today()
    return (today - timedelta(days=days)).strftime("%Y%m%d"), today.strftime("%Y%m%d")


def parse_any_datetime(text: Any):
    """'2026-10-09 12:30[:00]' · '2026.10.09' · '26.10.09' · '20261009' · '26/10/09' → datetime (정렬용). 못 읽으면 None"""
    s = str(text or "").strip()
    if not s:
        return None
    from datetime import datetime as _dt
    m = re.search(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?", s)
    if m:
        try:
            return _dt(int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4) or 0), int(m.group(5) or 0), int(m.group(6) or 0))
        except Exception:
            return None
    m = re.search(r"(?<!\d)(\d{2})[./-](\d{1,2})[./-](\d{1,2})(?!\d)", s)
    if m:
        try:
            return _dt(2000 + int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except Exception:
            return None
    m = re.search(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)", s)
    if m:
        try:
            return _dt(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except Exception:
            return None
    return None



# ───────────────────────── 평단가 (2026-10-13) ─────────────────────────
PYEONG_M2 = 3.3058
PYEONG_COLUMNS = ["기준면적(평)", "평단가(만원/평)", "평단가 기준"]
_TRADE_ORDER = {"매매": 0, "전세": 1, "월세": 2, "단기": 3, "단기임대": 3}


def _py_num(v):
    """'26,000' / 900000000 / 55.1 / None → float 또는 None"""
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return None
    try:
        return float(m.group(0))
    except Exception:
        return None


def _py_pair(v):
    """'26,000 / 100' → (26000.0, 100.0) / '55.1 / 16.67' → (55.1, 16.67) / 그 외 (num, None)"""
    if v is None:
        return (None, None)
    s = str(v)
    if "/" in s:
        a, b = s.split("/", 1)
        return (_py_num(a), _py_num(b))
    return (_py_num(s), None)


def pyeong_info(site: str, r: dict):
    """(거래, 금액(만원), 기준면적(평), 면적출처) — 금액/면적을 모르면 None"""
    site = (site or "").lower()
    trade = ""
    price = None
    area_py = None
    src = ""
    if site in ("daangn", "당근"):
        trade = str(r.get("거래유형") or "")
        if trade == "매매":
            price = _py_num(r.get("매매가"))
        elif trade == "전세":
            price = _py_num(r.get("전세금"))
        else:
            price = _py_num(r.get("월세"))
            if not trade:
                trade = "월세"
        m2 = _py_num(r.get("면적(㎡)") if r.get("면적(㎡)") not in (None, "") else r.get("면적"))
        if m2 and m2 > 0:
            area_py = m2 / PYEONG_M2
            src = "면적"
    elif site in ("naver", "네이버"):
        trade = str(r.get("거래방식") or "")
        if trade in ("매매", "전세"):
            price = _py_num(r.get("매매/전세금"))
        else:
            price = _py_num(r.get("월세"))
            if not trade:
                trade = "월세"
        if price is not None:
            price = price / 10000.0   # 네이버는 원 단위
        m2 = _py_num(r.get("공급/계약/대지"))
        src = "공급/계약면적"
        if not m2 or m2 <= 0:
            m2 = _py_num(r.get("전용/연"))
            src = "전용면적"
        if m2 and m2 > 0:
            area_py = m2 / PYEONG_M2
    else:  # 온하우스: '매매' / '전세' / '월세'("보증금 / 월세"), '전용면적(㎡/P)'("㎡ / 평")
        if str(r.get("매매") or "").strip():
            trade, price = "매매", _py_num(r.get("매매"))
        elif str(r.get("전세") or "").strip():
            trade, price = "전세", _py_num(r.get("전세"))
        elif str(r.get("월세") or "").strip():
            trade = "월세"
            _dep, rent = _py_pair(r.get("월세"))
            price = rent if rent is not None else _dep
        m2, py = _py_pair(r.get("전용면적(㎡/P)"))
        if py and py > 0:
            area_py = py
        elif m2 and m2 > 0:
            area_py = m2 / PYEONG_M2
        src = "전용면적"
    return trade, price, area_py, src


def add_pyeong_price(rows, site: str):
    """dict 행마다 기준면적(평)·평단가(만원/평)·평단가 기준 열을 붙인다. 면적을 모르거나 0이면 '확인불가'."""
    for r in rows:
        if not isinstance(r, dict):
            continue
        trade, price, area_py, src = pyeong_info(site, r)
        if area_py and area_py > 0 and price is not None and price > 0:
            val = round(price / area_py, 1)
            r["기준면적(평)"] = round(area_py, 2)
            r["평단가(만원/평)"] = val
            r["평단가 기준"] = f"{trade} {price:,.0f}만원 ÷ {area_py:.2f}평 ({src})"
        else:
            r["기준면적(평)"] = round(area_py, 2) if area_py else ""
            r["평단가(만원/평)"] = "확인불가"
            r["평단가 기준"] = "면적 없음(0)" if not area_py else f"{trade or '거래'} 금액 없음"
    return rows


def sort_by_pyeong_price(rows):
    """거래(매매→전세→월세) 안에서 평단가 낮은 순. '확인불가'는 맨 뒤."""
    def k(item):
        i, r = item
        trade = str((r or {}).get("거래유형") or (r or {}).get("거래방식") or "")
        if not trade:
            t, _p, _a, _s = pyeong_info("onhouse", r or {})
            trade = t
        v = (r or {}).get("평단가(만원/평)")
        try:
            fv = float(v)
            return (_TRADE_ORDER.get(trade, 9), 0, fv, i)
        except Exception:
            return (_TRADE_ORDER.get(trade, 9), 1, 0.0, i)
    return [r for _, r in sorted(enumerate(rows), key=k)]


def sort_newest_first(rows, keys):
    """dict 행 목록을 keys 중 처음 값이 있는 날짜로 최신순 정렬 (날짜 없는 행은 뒤로, 같은 날짜는 원래 순서)"""
    def k(item):
        i, r = item
        d = None
        for key in keys:
            d = parse_any_datetime((r or {}).get(key)) if isinstance(r, dict) else None
            if d:
                break
        return (0, -d.timestamp(), i) if d else (1, 0, i)
    return [r for _, r in sorted(enumerate(rows), key=k)]


def sort_rows_newest_first_by_index(rows, idx):
    """list 행 목록(네이버)을 idx 열의 날짜로 최신순 정렬"""
    def k(item):
        i, r = item
        d = parse_any_datetime(r[idx]) if isinstance(r, (list, tuple)) and len(r) > idx else None
        return (0, -d.timestamp(), i) if d else (1, 0, i)
    return [r for _, r in sorted(enumerate(rows), key=k)]


def parse_date_ymd(text: Any):
    s = str(text or "").strip()
    if not s:
        return None
    m = re.search(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})", s)
    if not m:
        return None
    try:
        from datetime import date

        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except Exception:
        return None


def description_text_from_naver_row(row: Any, *, naver_row_has_hosu_column: bool = False) -> str:
    """네이버 extract_detail_v2 행에서 상세 텍스트 필터용 문자열(간략설명+설명)을 만든다.
    naver_row_has_hosu_column=True: 호수가 인덱스 11에 삽입된 행(기본 60열 + 호수 = 61열 이상이면 간략/설명 인덱스 43,44).
    """
    if isinstance(row, (list, tuple)):
        if naver_row_has_hosu_column:
            if len(row) >= 61:
                a = str(row[43]) if len(row) > 43 else ""
                b = str(row[44]) if len(row) > 44 else ""
            else:
                a = str(row[42]) if len(row) > 42 else ""
                b = str(row[43]) if len(row) > 43 else ""
        else:
            a = str(row[42]) if len(row) > 42 else ""
            b = str(row[43]) if len(row) > 43 else ""
        return "\n".join(x for x in (a, b) if x)
    if isinstance(row, dict):
        a = str(row.get("간략설명", "") or "")
        b = str(row.get("설명", "") or "")
        return "\n".join(x for x in (a, b) if x)
    return ""


class RangeInput(QWidget):
    """숫자 직접 입력식 범위 필터 (최소 ~ 최대).

    - 최소값 빈칸 = 0, 최대값 빈칸 = 제한없음.
    - unit_toggle=(다른단위, 배율)을 주면 단위 변환 버튼이 생긴다.
      배율은 "기본단위 → 다른단위" 곱셈 계수이며 values()는 항상 기본단위로 반환.
    - DualSlider와 동일하게 values() -> (lo, hi), set_enabled() 인터페이스 제공.
    """

    OPEN_MAX = 10**9

    def __init__(
        self,
        title: str,
        unit: str = "",
        default_min_text: str = "",
        unit_toggle: Optional[tuple] = None,
        parent: Optional[QWidget] = None,
    ):
        super().__init__(parent)
        self.setObjectName("rangeCard")
        self._base_unit = unit
        self._alt = unit_toggle  # (다른단위, 배율) 또는 None
        self._in_alt = False
        row = QHBoxLayout(self)
        row.setContentsMargins(10, 6, 10, 6)
        row.setSpacing(6)
        self.title_label = QLabel(title)
        self.title_label.setStyleSheet("font-weight:600;")
        self.title_label.setFixedWidth(80)  # 라벨 폭 통일 → 입력칸 세로 정렬 맞춤
        self.edt_min = QLineEdit()
        self.edt_max = QLineEdit()
        for e in (self.edt_min, self.edt_max):
            e.setValidator(QIntValidator(0, 2_000_000_000, self))
            e.setFixedWidth(96)
            e.setAlignment(Qt.AlignRight)
        self.edt_min.setPlaceholderText("0")
        self.edt_max.setPlaceholderText("제한없음")
        self._default_min_text = str(default_min_text or "")
        if default_min_text:
            self.edt_min.setText(str(default_min_text))
        self.sep_label = QLabel("~")
        self.unit_label = QLabel(unit)
        self.unit_label.setStyleSheet("color:#64748b;")
        row.addWidget(self.title_label)
        row.addWidget(self.edt_min)
        row.addWidget(self.sep_label)
        row.addWidget(self.edt_max)
        row.addWidget(self.unit_label)
        self.btn_unit: Optional[QPushButton] = None
        if self._alt:
            self.btn_unit = QPushButton(f"{self._alt[0]}로 보기")
            self.btn_unit.setObjectName("rangeResetBtn")
            self.btn_unit.clicked.connect(self._toggle_unit)
            row.addWidget(self.btn_unit)
        row.addStretch(1)

    @staticmethod
    def _num(edit: QLineEdit) -> Optional[int]:
        t = edit.text().strip().replace(",", "")
        if not t:
            return None
        try:
            return int(t)
        except ValueError:
            return None

    def values(self) -> tuple:
        lo = self._num(self.edt_min) or 0
        hi = self._num(self.edt_max)
        hi = self.OPEN_MAX if hi is None else hi
        if self._alt and self._in_alt:  # 변환 단위로 입력 중이면 기본단위로 환산
            factor = self._alt[1]
            lo = int(round(lo / factor))
            if hi != self.OPEN_MAX:
                hi = int(round(hi / factor))
        if lo > hi:
            lo, hi = hi, lo
        return lo, hi

    def _toggle_unit(self):
        alt_unit, factor = self._alt
        conv = (lambda v: v * factor) if not self._in_alt else (lambda v: v / factor)
        for e in (self.edt_min, self.edt_max):
            v = self._num(e)
            if v is not None:
                e.setText(str(int(round(conv(v)))))
        self._in_alt = not self._in_alt
        self.unit_label.setText(alt_unit if self._in_alt else self._base_unit)
        self.btn_unit.setText(
            f"{self._base_unit}로 보기" if self._in_alt else f"{alt_unit}로 보기"
        )

    def reset(self) -> None:
        """처음 상태로: 최소=기본값(빈칸 또는 default_min_text), 최대=제한없음, 단위는 기본 단위."""
        if self._alt and self._in_alt:
            self._in_alt = False
            self.unit_label.setText(self._base_unit)
            if self.btn_unit is not None:
                self.btn_unit.setText(f"{self._alt[0]}로 보기")
        self.edt_min.setText(self._default_min_text)
        self.edt_max.setText("")

    def set_enabled(self, enabled: bool):
        for w in (self.title_label, self.edt_min, self.edt_max, self.sep_label, self.unit_label):
            w.setEnabled(enabled)
        if self.btn_unit is not None:
            self.btn_unit.setEnabled(enabled)


class MultiSelectCombo(QComboBox):
    """체크박스로 여러 항목을 동시에 선택할 수 있는 콤보박스.

    - 첫 항목은 항상 "전체". 전체를 체크하면 나머지가 해제되고,
      개별 항목을 체크하면 전체가 해제된다. 아무것도 선택하지 않은 상태는 "전체"로 유지.
    - 각 항목은 (표시 라벨, 내부 데이터) 쌍으로 등록한다.
    - checked_data(): 체크된 항목들의 내부 데이터 목록. 빈 목록이면 "전체"를 뜻한다.
    """

    selection_changed = Signal()
    ALL_LABEL = "전체"

    def __init__(self, parent: Optional[QWidget] = None):
        super().__init__(parent)
        self.setEditable(True)
        self.lineEdit().setReadOnly(True)
        self.setInsertPolicy(QComboBox.NoInsert)
        self.setModel(QStandardItemModel(self))
        self.view().viewport().installEventFilter(self)
        self.lineEdit().installEventFilter(self)
        self._updating = False
        self.model().itemChanged.connect(self._on_item_changed)
        self.set_entries([])

    # ── 항목 관리 ────────────────────────────────────────────────
    def set_entries(self, entries: List[tuple]) -> None:
        """entries: [(label, data), ...]. 기존 선택은 초기화("전체")된다."""
        self._updating = True
        m: QStandardItemModel = self.model()
        m.clear()
        all_item = QStandardItem(self.ALL_LABEL)
        all_item.setFlags(Qt.ItemIsEnabled | Qt.ItemIsUserCheckable)
        all_item.setData(None, Qt.UserRole)
        all_item.setCheckState(Qt.Checked)
        m.appendRow(all_item)
        for label, data in entries:
            it = QStandardItem(str(label))
            it.setFlags(Qt.ItemIsEnabled | Qt.ItemIsUserCheckable)
            it.setCheckState(Qt.Unchecked)
            it.setData(data, Qt.UserRole)
            m.appendRow(it)
        self._updating = False
        self._sync_text()
        self.selection_changed.emit()

    def set_checked_labels(self, labels: List[str]) -> None:
        """라벨 목록으로 체크 상태 지정 (없는 라벨은 무시)."""
        want = set(labels or [])
        self._updating = True
        m: QStandardItemModel = self.model()
        any_checked = False
        for r in range(1, m.rowCount()):
            it = m.item(r)
            on = it.text() in want
            it.setCheckState(Qt.Checked if on else Qt.Unchecked)
            any_checked = any_checked or on
        m.item(0).setCheckState(Qt.Unchecked if any_checked else Qt.Checked)
        self._updating = False
        self._sync_text()
        self.selection_changed.emit()

    # ── 선택 상태 조회 ───────────────────────────────────────────
    def checked_labels(self) -> List[str]:
        m: QStandardItemModel = self.model()
        return [
            m.item(r).text()
            for r in range(1, m.rowCount())
            if m.item(r).checkState() == Qt.Checked
        ]

    def checked_data(self) -> List[Any]:
        m: QStandardItemModel = self.model()
        return [
            m.item(r).data(Qt.UserRole)
            for r in range(1, m.rowCount())
            if m.item(r).checkState() == Qt.Checked
        ]

    def is_all(self) -> bool:
        return not self.checked_labels()

    # ── 내부 동작 ───────────────────────────────────────────────
    def _on_item_changed(self, item: QStandardItem) -> None:
        if self._updating:
            return
        self._updating = True
        m: QStandardItemModel = self.model()
        if item.row() == 0:
            if item.checkState() == Qt.Checked:
                for r in range(1, m.rowCount()):
                    m.item(r).setCheckState(Qt.Unchecked)
            elif not any(
                m.item(r).checkState() == Qt.Checked for r in range(1, m.rowCount())
            ):
                item.setCheckState(Qt.Checked)  # 전체 해제 불가(빈 선택 방지)
        else:
            if item.checkState() == Qt.Checked:
                m.item(0).setCheckState(Qt.Unchecked)
            elif not any(
                m.item(r).checkState() == Qt.Checked for r in range(1, m.rowCount())
            ):
                m.item(0).setCheckState(Qt.Checked)
        self._updating = False
        self._sync_text()
        self.selection_changed.emit()

    def _sync_text(self) -> None:
        labels = self.checked_labels()
        if not labels:
            text = self.ALL_LABEL
        elif len(labels) <= 2:
            text = ", ".join(labels)
        else:
            text = f"{labels[0]} 외 {len(labels) - 1}곳"
        self.lineEdit().setText(text)
        self.lineEdit().setCursorPosition(0)

    def eventFilter(self, obj, ev):
        # 팝업 안에서 클릭 시 체크만 토글하고 팝업은 닫지 않는다.
        if obj is self.view().viewport() and ev.type() == QEvent.MouseButtonRelease:
            idx = self.view().indexAt(ev.position().toPoint())
            it = self.model().itemFromIndex(idx)
            if it is not None:
                it.setCheckState(
                    Qt.Unchecked if it.checkState() == Qt.Checked else Qt.Checked
                )
            return True
        # 읽기전용 입력칸 클릭으로도 팝업 열기
        if obj is self.lineEdit() and ev.type() == QEvent.MouseButtonRelease:
            self.showPopup()
            return True
        return super().eventFilter(obj, ev)


def wrap_in_scroll(inner: QWidget) -> QScrollArea:
    """탭/섹션 콘텐츠를 세로 스크롤 영역으로 감싼다.

    창이 작아 아래쪽 섹션(가격/면적/등록일 등)이 잘리면 스크롤로 볼 수 있게 한다.
    setWidgetResizable(True) 라 가로 폭은 뷰포트에 맞춰지고 세로 스크롤만 필요 시 생긴다.
    """
    sa = QScrollArea()
    sa.setObjectName("sectionScroll")
    sa.setWidgetResizable(True)
    sa.setFrameShape(QFrame.NoFrame)
    sa.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
    sa.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
    sa.setWidget(inner)
    # 스크롤 영역/뷰포트 배경을 투명하게 두어 탭 pane 배경이 그대로 보이게 한다.
    sa.setStyleSheet(
        "#sectionScroll, #sectionScroll > QWidget > QWidget { background: transparent; }"
    )
    return sa


class _RangeBar(QWidget):
    """points 인덱스에 스냅되는 2-핸들 드래그 슬라이더 (커스텀 페인팅)."""

    MARGIN_X = 16
    TRACK_Y = 34
    TRACK_H = 8
    HANDLE_R = 9

    def __init__(self, n_points: int, label_fn, on_change, parent: QWidget | None = None):
        super().__init__(parent)
        self._n = max(1, n_points)
        self._label_fn = label_fn
        self._on_change = on_change
        self._lo = 0
        self._hi = self._n - 1
        self._active: Optional[str] = None
        self._hover: Optional[str] = None
        self._enabled = True
        self.setMinimumHeight(52)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.setMouseTracking(True)
        self.setCursor(Qt.PointingHandCursor)

    # --- 좌표 <-> 인덱스 변환 ---
    def _usable(self) -> float:
        return max(1.0, self.width() - 2 * self.MARGIN_X)

    def _x_of(self, idx: int) -> float:
        if self._n <= 1:
            return float(self.MARGIN_X)
        return self.MARGIN_X + self._usable() * (idx / (self._n - 1))

    def _idx_of(self, x: float) -> int:
        if self._n <= 1:
            return 0
        r = (x - self.MARGIN_X) / self._usable()
        return max(0, min(self._n - 1, round(r * (self._n - 1))))

    def _handle_at(self, x: float) -> str:
        d_lo = abs(x - self._x_of(self._lo))
        d_hi = abs(x - self._x_of(self._hi))
        if abs(d_lo - d_hi) < 1.0:
            return "lo" if x <= self._x_of(self._lo) else "hi"
        return "lo" if d_lo < d_hi else "hi"

    # --- 상태 변경 ---
    def set_range(self, lo: int, hi: int):
        self._lo = max(0, min(self._n - 1, lo))
        self._hi = max(0, min(self._n - 1, hi))
        if self._lo > self._hi:
            self._lo, self._hi = self._hi, self._lo
        self.update()
        if self._on_change:
            self._on_change()

    def set_bar_enabled(self, enabled: bool):
        self._enabled = enabled
        self.setCursor(Qt.PointingHandCursor if enabled else Qt.ArrowCursor)
        self.update()

    def _move_active_to(self, x: float):
        idx = self._idx_of(x)
        if self._active == "lo":
            self._lo = min(idx, self._hi)
        elif self._active == "hi":
            self._hi = max(idx, self._lo)
        self.update()
        if self._on_change:
            self._on_change()

    # --- 마우스 ---
    def mousePressEvent(self, e):
        if not self._enabled:
            return
        x = e.position().x()
        self._active = self._handle_at(x)
        self._move_active_to(x)

    def mouseMoveEvent(self, e):
        x = e.position().x()
        if self._active and self._enabled:
            self._move_active_to(x)
            return
        h = self._handle_at(x) if self._enabled else None
        if h != self._hover:
            self._hover = h
            self.update()

    def mouseReleaseEvent(self, e):
        self._active = None

    def leaveEvent(self, e):
        if self._hover is not None:
            self._hover = None
            self.update()

    # --- 페인팅 ---
    def paintEvent(self, e):
        p = QPainter(self)
        p.setRenderHint(QPainter.Antialiasing, True)

        x0 = self._x_of(0)
        x1 = self._x_of(self._n - 1)
        x_lo = self._x_of(self._lo)
        x_hi = self._x_of(self._hi)
        ty = self.TRACK_Y
        th = self.TRACK_H

        groove = QColor("#dbe3ef") if self._enabled else QColor("#e6ebf2")
        fill = QColor("#5ba0e7") if self._enabled else QColor("#c3d5ea")
        handle = QColor("#4f46e5") if self._enabled else QColor("#b7bce8")

        # 전체 트랙
        p.setPen(Qt.NoPen)
        p.setBrush(groove)
        p.drawRoundedRect(x0, ty - th / 2, x1 - x0, th, th / 2, th / 2)
        # 선택 구간
        p.setBrush(fill)
        p.drawRoundedRect(x_lo, ty - th / 2, max(1.0, x_hi - x_lo), th, th / 2, th / 2)

        # 핸들 위 값 라벨
        font = QFont(self.font())
        font.setBold(True)
        font.setPointSizeF(max(8.0, font.pointSizeF() - 0.5))
        p.setFont(font)
        fm = QFontMetrics(font)
        label_color = QColor("#1f3b5b") if self._enabled else QColor("#94a3b8")

        for idx, x in ((self._lo, x_lo), (self._hi, x_hi)):
            txt = self._label_fn(idx)
            tw = fm.horizontalAdvance(txt)
            tx = min(max(x - tw / 2, 2), self.width() - tw - 2)
            p.setPen(label_color)
            p.drawText(int(tx), 16, txt)

        # 핸들
        for x in (x_lo, x_hi):
            r = self.HANDLE_R
            p.setPen(QPen(QColor("#eef2ff"), 2))
            p.setBrush(handle)
            p.drawEllipse(int(x - r), int(ty - r), int(2 * r), int(2 * r))
        p.end()


class DualSlider(QWidget):
    # 상단 핸들이 맨 우측(마지막 눈금)일 때 values()의 상한으로 쓰는 값 = 상방 무제한.
    # 모든 슬라이더의 max_value보다 크므로 크롤러/필터에서 사실상 '제한 없음'으로 동작한다.
    OPEN_MAX = 10 ** 9

    def __init__(self, title: str, min_value: int, max_value: int, step: int = 1, parent: QWidget | None = None):
        super().__init__(parent)
        self.setObjectName("sliderCard")
        self.title_text = title
        self.min_value = min_value
        self.max_value = max_value
        self.step = max(1, step)
        self.points = self._build_points(title, min_value, max_value, self.step)
        self._is_money = ("만" in title) or ("매매" in title) or ("보증금" in title) or ("전세" in title) or ("기준가" in title) or ("월세" in title)
        self._is_year = "년도" in title

        root = QVBoxLayout(self)
        root.setContentsMargins(12, 10, 12, 10)
        root.setSpacing(6)

        row = QHBoxLayout()
        self.title_label = QLabel(title)
        self.value_label = QLabel()
        self.value_label.setStyleSheet("color:#64748b;")
        self.btn_reset = QPushButton("초기화")
        self.btn_reset.setObjectName("rangeResetBtn")
        self.btn_reset.clicked.connect(self._reset)
        row.addWidget(self.title_label)
        row.addStretch(1)
        row.addWidget(self.value_label)
        row.addWidget(self.btn_reset)
        root.addLayout(row)

        self.bar = _RangeBar(len(self.points), self._handle_text, self._on_bar_change)
        root.addWidget(self.bar)
        self._sync_ui()

    def _fmt_val(self, v: int) -> str:
        if self._is_year:
            return str(v)
        return self._format_unit(v, self._is_money)

    def _handle_text(self, idx: int) -> str:
        idx = max(0, min(len(self.points) - 1, idx))
        txt = self._fmt_val(self.points[idx])
        # 맨 우측 눈금은 '이상'(상방 무제한)을 뜻하므로 '+'를 붙인다.
        if idx == len(self.points) - 1 and len(self.points) > 1:
            txt = f"{txt}+"
        return txt

    def _on_bar_change(self):
        self._sync_ui()

    def _build_points(self, title: str, min_value: int, max_value: int, step: int) -> List[int]:
        t = title.replace(" ", "")
        if "월세" in t:
            base = [min_value, 20, 30, 40, 50, 60, 70, 80, 90, 100, 200, 300, 400, 500, 1000, 2000, 3000]
        elif "보증금" in t:
            base = [min_value, 1000, 2000, 3000, 4000, 5000, 6000, 7000, 8000, 9000, 10000, 20000, 30000, 40000, 50000]
        elif "매매" in t or "전세" in t or "기준가" in t:
            base = [min_value, 5000, 10000, 20000, 30000, 40000, 50000, 60000, 70000, 80000, 90000, 100000, 200000, 300000, 400000, 500000, 1000000]
        elif "면적" in t or "평수" in t:
            base = [min_value, 10, 20, 30, 50, 70, 100, 150, 200, 231, 500]
        else:
            base = list(range(min_value, max_value + 1, step))

        pts = sorted(set([x for x in base if min_value <= x <= max_value] + [min_value, max_value]))
        return pts

    @staticmethod
    def _format_unit(v: int, is_money: bool) -> str:
        if not is_money:
            return f"{v:,}"
        if v >= 10000:
            if v % 10000 == 0:
                return f"{v // 10000}억"
            return f"{v:,}"
        if v >= 1000 and v % 1000 == 0:
            return f"{v // 1000}천"
        if v >= 100 and v % 100 == 0:
            return f"{v // 100}백"
        return f"{v:,}"

    def _reset(self):
        self.bar.set_range(0, len(self.points) - 1)
        self._sync_ui()

    def _sync_ui(self):
        last = len(self.points) - 1
        lo_idx, hi_idx = self.bar._lo, self.bar._hi
        open_hi = hi_idx >= last  # 맨 우측 = 상방 무제한
        lo_txt = self._fmt_val(self.points[lo_idx])
        hi_txt = self._fmt_val(self.points[hi_idx])
        if lo_idx <= 0 and open_hi:
            text = "전체"
        elif open_hi:
            text = f"{lo_txt} 이상"
        elif lo_idx <= 0:
            text = f"~{hi_txt}"
        else:
            text = f"{lo_txt} ~ {hi_txt}"
        self.value_label.setText(text)

    def values(self) -> tuple[int, int]:
        """(하한, 상한). 상단 핸들이 맨 우측이면 상한은 OPEN_MAX(상방 무제한)."""
        lo = self.points[self.bar._lo]
        if self.bar._hi >= len(self.points) - 1:
            return lo, self.OPEN_MAX
        return lo, self.points[self.bar._hi]

    def set_enabled(self, enabled: bool):
        self.title_label.setEnabled(enabled)
        self.value_label.setEnabled(enabled)
        self.btn_reset.setEnabled(enabled)
        self.bar.set_bar_enabled(enabled)


# 건축물용도 다중선택 옵션 (데이터 값에 부분일치로 매칭)
# 건축물용도 선택지. 당근(realty.daangn.com buildingUsage 한글 라벨)과 네이버(buildingUse 문자열) 값을
# 공백/구분자 제거 후 부분일치로 비교하므로 "제2종 근린생활시설"도 "제2종근린생활시설"에 걸린다.
BUILDING_USE_OPTIONS = [
    "단독주택", "공동주택", "다가구주택", "다세대주택", "연립주택", "아파트",
    "오피스텔", "제1종근린생활시설", "제2종근린생활시설", "업무시설", "판매시설", "숙박시설",
    "의료시설", "교육연구시설", "노유자시설", "문화및집회시설", "종교시설", "운동시설",
    "위락시설", "공장", "창고시설", "자동차관련시설", "동물및식물관련시설", "운수시설",
]


def _leading_int(s: Any) -> Optional[int]:
    """문자열에서 맨 앞 정수를 뽑는다. 예: '3/2' -> 3, '5개' -> 5."""
    m = re.search(r"\d+", str(s or "").replace(",", ""))
    return int(m.group()) if m else None


def _year4(s: Any) -> Optional[int]:
    """문자열에서 4자리 연도를 뽑는다. 예: '2001.05.20' -> 2001."""
    m = re.search(r"(19|20)\d{2}", str(s or ""))
    return int(m.group()) if m else None


class DetailFilters:
    """방수/층수/준공년도/건축물용도 후처리 필터.

    - rooms/floors/years: (min, max) 튜플 또는 None(미적용)
    - uses: 선택된 건축물용도 라벨 리스트 (빈 리스트=미적용)
    모든 필드는 데이터 값이 비어 있으면 관대하게 통과한다(알 수 있는 값만 걸러냄).
    건축물용도는 값이 있으면 선택 라벨과 부분일치해야 통과.
    """

    def __init__(self, rooms=None, floors=None, years=None, uses=None, uses_strict: bool = False):
        self.rooms = rooms
        self.floors = floors
        self.years = years
        self.uses = [str(u) for u in (uses or []) if str(u).strip()]
        self.uses_strict = bool(uses_strict)   # True: 용도를 체크했는데 매물에 용도 값이 없으면 제외 (네이버는 lawUsage 가 거의 항상 있음)

    @property
    def active(self) -> bool:
        return bool(self.rooms or self.floors or self.years or self.uses)

    @staticmethod
    def _num_ok(bound, s) -> bool:
        if not bound:
            return True
        v = _leading_int(s)
        if v is None:
            return True
        lo, hi = bound
        return lo <= v <= hi

    def _year_ok(self, s) -> bool:
        if not self.years:
            return True
        y = _year4(s)
        if y is None:
            return True
        lo, hi = self.years
        return lo <= y <= hi

    def _use_ok(self, s) -> bool:
        if not self.uses:
            return True
        val = re.sub(r"[\s/·,()\[\]]", "", str(s or ""))
        if not val:
            return not self.uses_strict  # 용도 값이 없으면 관대하게 통과(당근 목록에 값이 비는 경우 대비) · strict 면 제외
        for u in self.uses:
            uu = re.sub(r"[\s/·,()\[\]]", "", u)
            if uu and (uu in val or val in uu):
                return True
        return False

    def _floor_ok(self, s) -> bool:
        """층수 필터. 반지하/지하(B1 등)는 0층으로 취급한다."""
        if not self.floors:
            return True
        t = str(s or "")
        if "반지하" in t or "지하" in t or t.strip().upper().startswith("B"):
            v = 0
        else:
            v = _leading_int(t)
        if v is None:
            return True
        lo, hi = self.floors
        return lo <= v <= hi

    def passes(self, *, rooms: str = "", floor: str = "", approve: str = "", use: str = "") -> bool:
        return (
            self._num_ok(self.rooms, rooms)
            and self._floor_ok(floor)
            and self._year_ok(approve)
            and self._use_ok(use)
        )


def naver_detail_values(row: Any, has_hosu: bool = False) -> tuple[str, str, str, str]:
    """네이버 상세 행에서 (방수, 해당층, 사용승인일, 건축물용도) 문자열을 뽑는다.

    has_hosu=True: '호수' 컬럼이 인덱스 11에 삽입된 레이아웃(app_v2)이라 11 이후 인덱스가 +1 밀린다.
    """
    shift = 1 if has_hosu else 0

    def _get(i: int) -> str:
        j = i + shift if i >= 11 else i
        if isinstance(row, (list, tuple)) and len(row) > j and row[j] is not None:
            return str(row[j]).strip()
        return ""
    # 기본 인덱스: 34 방수, 17 해당층, 49 사용승인일, 39 건축물용도
    return _get(34), _get(17), _get(49), _get(39)


def _dict_find(item: dict, includes, excludes=()) -> str:
    """dict에서 키가 includes 문자열을 모두 포함하고 excludes는 포함하지 않는 첫 값(문자열)을 반환."""
    for k, v in item.items():
        ks = str(k)
        if all(i in ks for i in includes) and not any(e in ks for e in excludes):
            sv = str(v).strip()
            if sv:
                return sv
    return ""


def daangn_detail_values(item: Any) -> tuple[str, str, str, str]:
    """당근 매물 dict에서 (방수, 층수, 사용승인일, 건축물용도) 문자열을 best-effort로 뽑는다.

    당근 상세 라벨(dt/dd)이 그대로 키가 되므로 유동적 → 부분 키매칭 사용.
    """
    if not isinstance(item, dict):
        return "", "", "", ""
    rooms = _dict_find(item, ["방"], excludes=["방향", "방문"])
    floor = _dict_find(item, ["층"])
    approve = _dict_find(item, ["사용승인"]) or _dict_find(item, ["준공"])
    use = _dict_find(item, ["건축물용도"]) or _dict_find(item, ["용도"], excludes=["용도지역"])
    return rooms, floor, approve, use


class UseCheckGroup(QGroupBox):
    """건축물용도 다중선택 체크박스 그룹. selected()로 체크된 라벨 리스트 반환."""

    def __init__(self, title: str = "건축물용도", options: Optional[List[str]] = None, columns: int = 4, parent: QWidget | None = None):
        super().__init__(title, parent)
        self._checks: dict[str, QCheckBox] = {}
        opts = options if options is not None else BUILDING_USE_OPTIONS
        grid = QGridLayout(self)
        grid.setContentsMargins(12, 14, 12, 12)
        grid.setHorizontalSpacing(10)
        grid.setVerticalSpacing(6)
        for i, name in enumerate(opts):
            cb = QCheckBox(name)
            self._checks[name] = cb
            grid.addWidget(cb, i // columns, i % columns)

    def selected(self) -> List[str]:
        return [n for n, cb in self._checks.items() if cb.isChecked()]

    def set_selected(self, names: List[str]) -> None:
        wanted = set(names or [])
        for n, cb in self._checks.items():
            cb.setChecked(n in wanted)

    def set_enabled(self, enabled: bool) -> None:
        for cb in self._checks.values():
            cb.setEnabled(enabled)


__all__ = [
    "BUILDING_USE_OPTIONS",
    "DetailFilters",
    "DualSlider",
    "MultiSelectCombo",
    "RangeInput",
    "UseCheckGroup",
    "daangn_detail_values",
    "naver_detail_values",
    "contains_any_keyword",
    "description_text_from_naver_row",
    "open_path_in_os",
    "parse_date_ymd",
    "parse_keywords_csv",
    "parse_saved_output_path_from_finish_message",
    "wrap_in_scroll",
]


# 통합 수집기(app_v1_site_main)와 같은 화면 스타일 — 온하우스 단독 수집기도 같은 모양으로 (2026-10-01)
APP_STYLESHEET = """
            QMainWindow {
                background: #1f2431;
            }
            QWidget#appRoot {
                background: qlineargradient(
                    x1: 0, y1: 0, x2: 1, y2: 1,
                    stop: 0 #30384a,
                    stop: 0.5 #242b3a,
                    stop: 1 #1c2230
                );
            }
            QWidget {
                color: #f8fafc;
                font-family: 'Segoe UI', 'Malgun Gothic', 'Yu Gothic UI', sans-serif;
                font-size: 12px;
            }
            QMessageBox {
                background: #ffffff;
            }
            QMessageBox QLabel {
                color: #111827;
                font-weight: 600;
            }
            QMessageBox QPushButton {
                background: #4f46e5;
                color: #ffffff;
                border: none;
                border-radius: 8px;
                padding: 7px 14px;
                min-width: 84px;
                font-weight: 700;
            }
            QMessageBox QPushButton:hover { background: #4338ca; }
            QMessageBox QPushButton:pressed { background: #3730a3; }

            QTabWidget::pane {
                border: 1px solid #7c8da8;
                border-radius: 12px;
                background: rgba(250, 252, 255, 0.97);
                top: -1px;
            }
            QTabBar::tab {
                background: #c9d3e2;
                color: #253248;
                border: 1px solid #9caec6;
                border-bottom: none;
                padding: 10px 18px;
                margin-right: 6px;
                border-top-left-radius: 10px;
                border-top-right-radius: 10px;
                font-weight: 600;
            }
            QTabBar::tab:hover { background: #b9c7da; color: #1e293b; }
            QTabBar::tab:selected { background: #4f46e5; color: white; border-color: #4f46e5; }

            QGroupBox {
                border: 1px solid #d7e0ec;
                border-radius: 14px;
                margin-top: 12px;
                padding-top: 16px;
                background: rgba(249, 251, 255, 0.98);
                font-weight: 700;
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 12px;
                padding: 0 6px;
                color: #334155;
                background: transparent;
            }

            QLabel { color: #f1f5f9; }
            QGroupBox QLabel { color: #0f172a; font-weight: 600; }
            QCheckBox { spacing: 6px; }
            QGroupBox QCheckBox { color: #111827; font-weight: 600; }
            QGroupBox QCheckBox[dimmed="true"] { color: #6b7280; font-weight: 700; }
            QCheckBox::indicator {
                width: 16px;
                height: 16px;
                border: 1px solid #a7b3c5;
                border-radius: 5px;
                background: #ffffff;
            }
            QCheckBox::indicator:checked {
                background: #4f46e5;
                border: 1px solid #4338ca;
            }
            QGroupBox QCheckBox[dimmed="true"]::indicator {
                border: 1px solid #9ca3af;
                background: #e5e7eb;
            }
            QGroupBox QCheckBox[dimmed="true"]::indicator:checked {
                border: 1px solid #6b7280;
                background: #6b7280;
            }

            QLineEdit, QComboBox, QDateEdit {
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                padding: 6px 8px;
                background: #ffffff;
                color: #0f172a;
                selection-background-color: #4f46e5;
            }
            QLineEdit::placeholder {
                color: #64748b;
            }
            /* 거래유형 미선택 등으로 비활성화된 입력칸은 어둡게 표시 */
            QLineEdit:disabled {
                background: #e5e7eb;
                color: #9ca3af;
                border-color: #d1d5db;
            }
            QGroupBox QLabel:disabled { color: #b0b8c4; }
            QPushButton:disabled {
                background: #e5e7eb;
                color: #9ca3af;
                border-color: #d1d5db;
            }
            QComboBox QAbstractItemView {
                color: #0f172a;
                background: #ffffff;
                selection-background-color: #c7d2fe;
                selection-color: #111827;
                border: 1px solid #cbd5e1;
            }
            QDateEdit::drop-down {
                subcontrol-origin: padding;
                subcontrol-position: top right;
                width: 22px;
                border-left: 1px solid #cbd5e1;
            }
            QDateEdit[dimmed="true"] {
                background: #e5e7eb;
                color: #94a3b8;
                border-color: #cbd5e1;
            }
            QLabel[dimmed="true"] {
                color: #94a3b8;
            }
            QCalendarWidget QWidget {
                background: #ffffff;
                color: #0f172a;
                selection-background-color: #4f46e5;
                selection-color: #ffffff;
            }
            QCalendarWidget QToolButton {
                color: #0f172a;
                background: #eef2ff;
                border: 1px solid #d4dce8;
                border-radius: 6px;
                font-weight: 700;
                padding: 4px 8px;
            }
            QCalendarWidget QToolButton:hover { background: #e0e7ff; }
            QCalendarWidget QSpinBox {
                color: #0f172a;
                background: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 6px;
                padding: 2px 4px;
            }
            QCalendarWidget QMenu {
                background: #ffffff;
                color: #0f172a;
                border: 1px solid #cbd5e1;
            }
            QCalendarWidget QMenu::item:selected {
                background: #c7d2fe;
                color: #111827;
            }
            QCalendarWidget QAbstractItemView:enabled {
                color: #0f172a;
                background: #ffffff;
                selection-background-color: #4f46e5;
                selection-color: #ffffff;
            }
            QLineEdit:hover, QComboBox:hover, QDateEdit:hover { border-color: #94a3b8; }
            QLineEdit:focus, QComboBox:focus, QDateEdit:focus { border-color: #4f46e5; }

            QPushButton {
                background: #4f46e5;
                color: white;
                border: none;
                border-radius: 10px;
                padding: 9px 16px;
                font-weight: 700;
            }
            QPushButton:hover { background: #4338ca; }
            QPushButton:pressed { background: #3730a3; }
            QPushButton:disabled { background: #a5b4fc; color: #eef2ff; }
            QPushButton#rangePresetBtn {
                background: #f8fbff;
                color: #2f3f56;
                border: 1px solid #b8d2ef;
                border-radius: 5px;
                padding: 5px 8px;
                min-width: 34px;
                font-weight: 700;
            }
            QPushButton#rangePresetBtn:hover {
                background: #eef6ff;
                border-color: #8fb8e6;
            }
            QPushButton#rangePresetBtn[inRange="true"] {
                background: #b7d8f7;
                color: #1f3b5b;
                border-color: #84b8e8;
            }
            QPushButton#rangePresetBtn[inRange="true"]:hover {
                background: #a7cef3;
                border-color: #6ea9e2;
            }
            QPushButton#rangePresetBtn:checked {
                background: #5ba0e7;
                color: #ffffff;
                border-color: #4a8fd6;
            }
            QPushButton#rangePresetBtn:disabled {
                background: #f1f5f9;
                color: #94a3b8;
                border-color: #d7e2ef;
            }
            QPushButton#rangeResetBtn {
                background: #9ca3af;
                color: #ffffff;
                border: none;
                border-radius: 5px;
                padding: 5px 10px;
                min-width: 48px;
            }
            QPushButton#rangeResetBtn:hover { background: #6b7280; }
            QPushButton#rangeResetBtn:pressed { background: #4b5563; }
            QPushButton#rangeResetBtn:disabled { background: #cbd5e1; color: #f8fafc; }

            QSlider::groove:horizontal {
                border: 0;
                height: 12px;
                background: #dbe3ef;
                border-radius: 6px;
            }
            QSlider::handle:horizontal {
                width: 22px;
                margin: -6px 0;
                border-radius: 11px;
                background: #4f46e5;
                border: 2px solid #eef2ff;
            }
            QSlider::handle:horizontal:hover { background: #4338ca; }

            QWidget#sliderCard {
                background: #e8eef6;
                border: 1px solid #d8e1ed;
                border-radius: 12px;
            }

            QHeaderView::section {
                background: #eaf0f7;
                color: #111827;
                padding: 8px;
                border: 1px solid #d4dce8;
                font-weight: 700;
            }
            QTableWidget {
                background: #ffffff;
                alternate-background-color: #f8fafc;
                color: #0f172a;
                gridline-color: #e2e8f0;
                border: 1px solid #d4dce8;
                border-radius: 10px;
                selection-background-color: #a5b4fc;
                selection-color: #111827;
            }

            QProgressBar {
                border: 1px solid #cbd5e1;
                border-radius: 7px;
                text-align: center;
                background: #e2e8f0;
                min-height: 16px;
            }
            QProgressBar::chunk {
                background: #4f46e5;
                border-radius: 6px;
            }
            """


# ───────────── 활동 기록용 수집 조건 요약 (2026-10-14) ─────────────
NAVER_TYPE_KO = {"JGC": "재건축", "IA01:IA02:IC01:IC02:IA04:IC03": "분양중/예정", "OBYG": "오피스텔분양권", "APT": "아파트",
                 "ABYG": "아파트분양권", "JGB": "재개발", "OPST": "오피스텔", "VL": "빌라/연립", "DDDGG": "단독/다가구",
                 "JWJT": "전원주택", "SGJT": "상가주택", "HOJT": "한옥주택", "OR": "원룸/투룸", "SG": "상가", "SMS": "사무실",
                 "GJCG": "공장/창고", "APTHGJ": "지식산업센터", "GM": "건물", "TJ": "토지"}
NAVER_TRADE_KO = {"A1": "매매", "B1": "전세", "B2": "월세", "B3": "단기"}
DAANGN_TRADE_KO = {"month": "월세", "buy": "매매", "borrow": "전세", "short": "단기"}
WRITER_KO = {"BROKER": "중개사", "DIRECT_USER": "직거래"}


def _daangn_sales_ko(code: str) -> str:
    try:
        from daangn_realty_crawler import DaangnRealtyCrawler
        return DaangnRealtyCrawler.SALES_TYPE_MAP.get(code, code)
    except Exception:
        return code


def describe_params(p, max_len: int = 380) -> str:
    """수집 파라미터 → 어드민 활동 기록에 남길 한 줄 조건 요약 (지역·유형·거래·기간·금액·세부·키워드·옵션).
    통합(당근/네이버/온하우스/피터팬)·온하우스 단독·예약 파라미터의 키를 모두 알아본다. 모르는 키는 무시."""
    if not isinstance(p, dict):
        return ""
    parts = []

    def num(v):
        if v is None or v == "":
            return ""
        try:
            f = float(v)
            return str(int(f)) if f == int(f) else f"{f:g}"
        except Exception:
            return str(v)

    def rng(label, a, b, unit="", lo=None, hi=None, div=1):
        if a is None and b is None:
            return
        try:
            a2 = None if a is None else float(a) / div
            b2 = None if b is None else float(b) / div
        except Exception:
            return
        a_open = a2 is None or a2 <= (lo if lo is not None else 0)
        b_open = b2 is None or (hi is not None and b2 >= hi)
        if a_open and b_open:
            return
        if b_open:
            parts.append(f"{label} {num(a2)}{unit} 이상")
        elif a_open:
            parts.append(f"{label} {num(b2)}{unit} 이하")
        else:
            parts.append(f"{label} {num(a2)}~{num(b2)}{unit}")

    if p.get("schedule_name"):
        parts.append(f"예약 '{p['schedule_name']}'")
    # 지역
    regions = p.get("region_names") or p.get("regions") or []
    if isinstance(regions, (list, tuple)) and regions:
        label = p.get("region_label") or p.get("region_for_file") or ""
        label = str(label).replace("_", " ").strip()
        first = str(regions[0])
        if label and label != "전체":
            parts.append(f"지역 {label} ({len(regions)}곳)")
        elif len(regions) == 1:
            parts.append(f"지역 {first}")
        else:
            parts.append(f"지역 {first} 외 {len(regions) - 1}곳")
    # 매물 유형
    kinds = []
    rt = p.get("realestate_type")
    if rt:
        kinds += [NAVER_TYPE_KO.get(c, c) for c in str(rt).split(":") if c]
    st = p.get("sales_type")
    if st:
        kinds += [_daangn_sales_ko(c) for c in str(st).split(",") if c]
    if p.get("room_type") and p.get("room_type") != "all":
        kinds.append(str(p["room_type"]))
    bt = p.get("building_types") or p.get("building_type")
    if bt:
        kinds += [str(x) for x in (bt if isinstance(bt, (list, tuple)) else [bt]) if x]
    if kinds:
        parts.append("유형 " + ", ".join(dict.fromkeys(kinds)))
    # 거래 방식
    trades = []
    tt = p.get("trade_type")
    if tt:
        for c in str(tt).replace(":", ",").split(","):
            c = c.strip()
            if c:
                trades.append(DAANGN_TRADE_KO.get(c, NAVER_TRADE_KO.get(c, c)))
    for key in ("trade_types", "contract_types"):
        v = p.get(key)
        if isinstance(v, (list, tuple)):
            trades += [str(x).strip("'\"") for x in v if x]
        elif isinstance(v, str) and v:
            trades += [x.strip().strip("'\"") for x in v.split(",") if x.strip()]
    if trades:
        parts.append("거래 " + "/".join(dict.fromkeys(trades)))
    # 기간
    if p.get("date_start") or p.get("date_end"):
        parts.append(f"등록 {p.get('date_start') or ''}~{p.get('date_end') or ''}")
    elif p.get("min_date"):
        parts.append(f"등록 {p['min_date']} 이후")
    elif p.get("period_label"):
        parts.append(f"{'등록일' if p.get('date_by') == 'reg' else '확인일'} {p['period_label']}")
    elif p.get("date_from") or p.get("date_to"):
        parts.append(f"{'등록일' if p.get('date_by') == 'reg' else '확인일'} {p.get('date_from') or ''}~{p.get('date_to') or ''}")
    elif p.get("registered_date_start") or p.get("registered_date_end"):
        parts.append(f"등록 {p.get('registered_date_start') or ''}~{p.get('registered_date_end') or ''}")
    # 금액 (만원) — 슬라이더 전체 범위면 생략
    rng("월세", p.get("monthly_min"), p.get("monthly_max"), "만", 0, 3000)
    rng("보증금", p.get("deposit_min"), p.get("deposit_max"), "만", 0, 50000, div=(10000 if p.get("site_name") == "피터팬" else 1))
    rng("매매", p.get("buy_min"), p.get("buy_max"), "만", 0, 1000000)
    rng("전세", p.get("borrow_min"), p.get("borrow_max"), "만", 0, 1000000)
    rng("면적", p.get("area_min"), p.get("area_max"), "㎡", 0, 200)
    rng("매매", p.get("min_deal"), p.get("max_deal"), "만")
    rng("보증금", p.get("min_warr"), p.get("max_warr"), "만")
    rng("월세", p.get("min_rent"), p.get("max_rent"), "만")
    rng("면적", p.get("min_area"), p.get("max_area"), "㎡")
    rng("보증금", p.get("refer_min"), p.get("refer_max"), "만")
    rng("매매/전세", p.get("price_min"), p.get("price_max"), "만", div=(10000 if p.get("site_name") == "피터팬" else 1))
    rng("월세", p.get("month_min"), p.get("month_max"), "만", div=(10000 if p.get("site_name") == "피터팬" else 1))
    rng("면적", p.get("size_min"), p.get("size_max"), "㎡")
    # 세부 조건
    for key, label, unit in (("detail_rooms", "방", "개"), ("post_rooms", "방", "개"), ("detail_floors", "층", "층"),
                             ("post_floors", "층", "층"), ("detail_years", "준공", "년"), ("post_years", "준공", "년")):
        v = p.get(key)
        if isinstance(v, (list, tuple)) and len(v) == 2:
            rng(label, v[0], v[1], unit)
    uses = p.get("detail_uses")
    if isinstance(uses, (list, tuple)) and uses:
        parts.append("용도 " + "/".join(str(u) for u in uses))
    kws = p.get("detail_keywords") or p.get("keywords")
    if isinstance(kws, (list, tuple)) and kws:
        parts.append("키워드 " + ",".join(str(k) for k in kws))
    elif isinstance(kws, str) and kws.strip():
        parts.append("키워드 " + kws.strip())
    # 옵션
    opts = []
    wt = p.get("writer_types")
    if isinstance(wt, (list, tuple)) and wt and len(set(wt)) < 2:
        opts.append(WRITER_KO.get(list(wt)[0], str(list(wt)[0])) + "만")
    if p.get("mode") == "share":
        opts.append("공동중개")
    if p.get("short_term"):
        opts.append("단기만")
    if p.get("hosu_enabled"):
        opts.append("호수")
    if p.get("contact"):
        opts.append("연락처")
    if p.get("sort_views"):
        opts.append("조회수순")
    if p.get("sort_pyeong"):
        opts.append("평단가순")
    if p.get("dedupe_same") or p.get("dedupe"):
        opts.append("중복제거")
    if p.get("max_pages"):
        opts.append(f"최대 {p['max_pages']}쪽")
    if opts:
        parts.append("옵션 " + "/".join(opts))
    s = " · ".join(parts)
    return s if len(s) <= max_len else s[: max_len - 1] + "…"
